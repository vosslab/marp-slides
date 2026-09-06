#!/usr/bin/env python3
"""Convert a trusted PPTX into experimental extended-Djot slide source."""

# Standard Library
import os
import re
import json
import pathlib
import argparse
import hashlib
import tempfile
import dataclasses

# PIP3 modules
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# Local modules
from tools import pptx_to_marp as pptx_common


WMF_HEADERS = (b"\xd7\xcd\xc6\x9a", b"\x01\x00\x09\x00\x00\x03", b"\x02\x00\x09\x00\x00\x03")
EMF_SIGNATURE_OFFSET = 40
EMF_SIGNATURE = b" EMF"
SUPPORTED_IMAGE_SUFFIXES = {".emf", ".gif", ".jpg", ".png", ".wmf"}
PRIMED_SEQUENCE = re.compile(
	r"([35])[′’']-([ACGTU][ACGTU|/,.]{2,}[ACGTU])-[′’']([35])"
)
PRIME_MARK = re.compile(r"([35])[′’']")
DNA_SEQUENCE = re.compile(r"(?<![A-Za-z0-9`])([ACGTU][ACGTU|/,.]{2,}[ACGTU])(?![A-Za-z0-9`])")


@dataclasses.dataclass(frozen=True)
class ConversionSummary:
	"""User-facing counts for one Djot conversion."""

	visible_slides: int
	editable_slides: int
	hidden_slides: int
	extracted_images: int
	review_slides: int
	output_path: pathlib.Path
	report_path: pathlib.Path


#============================================
def djot_text(text: str) -> str:
	"""Normalize short DNA notation into settled Djot authoring forms."""
	text = PRIMED_SEQUENCE.sub(r"\1&prime;-`\2`-\3&prime;", text)
	text = PRIME_MARK.sub(r"\1&prime;", text)
	return DNA_SEQUENCE.sub(r"`\1`", text)


#============================================
def validate_image_blob(blob: bytes, suffix: str) -> None:
	"""Validate one raster, WMF, or EMF image before writing it."""
	# ASVS 5.2.2 and 5.2.6: validate the declared type and bounded content.
	if suffix not in SUPPORTED_IMAGE_SUFFIXES:
		raise ValueError(f"unsupported PPTX image type: {suffix}")
	if len(blob) > pptx_common.MAX_MEMBER_BYTES:
		raise ValueError("PPTX image exceeds the per-image size limit")
	if suffix == ".wmf":
		# ASVS 5.2.2: accept only recognized WMF headers.
		if not blob.startswith(WMF_HEADERS):
			raise ValueError("invalid WMF image header")
		return
	if suffix == ".emf":
		# ASVS 5.2.2: validate an EMF record type and its required signature.
		if blob[:4] != b"\x01\x00\x00\x00" or blob[EMF_SIGNATURE_OFFSET:44] != EMF_SIGNATURE:
			raise ValueError("invalid EMF image header")
		return
	pptx_common.validate_image_blob(blob, suffix)


#============================================
def image_suffix(blob: bytes, declared_suffix: str) -> str:
	"""Correct LibreOffice's EMF-as-WMF metadata before validation."""
	if declared_suffix == ".wmf" and blob[EMF_SIGNATURE_OFFSET:44] == EMF_SIGNATURE:
		return ".emf"
	return declared_suffix


#============================================
def image_asset(
	shape: object,
	assets_dir: pathlib.Path,
	djot_root: pathlib.PurePosixPath,
	known_images: dict[str, str],
) -> pptx_common.ImageAsset:
	"""Extract one picture into a content-addressed source asset."""
	blob = shape.image.blob
	suffix = f".{shape.image.ext.lower()}"
	if suffix == ".jpeg":
		suffix = ".jpg"
	suffix = image_suffix(blob, suffix)
	validate_image_blob(blob, suffix)
	digest = hashlib.sha256(blob).hexdigest()
	asset_name = known_images.get(digest)
	if asset_name is None:
		asset_name = f"image_{len(known_images) + 1:03d}{suffix}"
		# ASVS 5.3.2: an archive name never selects the output destination.
		(assets_dir / asset_name).write_bytes(blob)
		known_images[digest] = asset_name
	djot_path = (djot_root / asset_name).as_posix()
	asset_number = int(re.search(r"\d+", asset_name).group())
	return pptx_common.ImageAsset(
		left=shape.left,
		top=shape.top,
		width=shape.width,
		height=shape.height,
		markdown_path=djot_path,
		alt_text=f"Slide image {asset_number}",
	)


#============================================
def shape_inventory(
	shape: object,
	title_id: int | None,
	assets_dir: pathlib.Path,
	djot_root: pathlib.PurePosixPath,
	known_images: dict[str, str],
) -> tuple[list[str], list[pptx_common.TextBlock], list[pptx_common.ImageAsset], list[str]]:
	"""Extract supported semantic objects recursively from one shape."""
	titles: list[str] = []
	text_blocks: list[pptx_common.TextBlock] = []
	images: list[pptx_common.ImageAsset] = []
	review_reasons: list[str] = []
	if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
		for child in shape.shapes:
			child_parts = shape_inventory(child, title_id, assets_dir, djot_root, known_images)
			titles.extend(child_parts[0])
			text_blocks.extend(child_parts[1])
			images.extend(child_parts[2])
			review_reasons.extend(child_parts[3])
		return titles, text_blocks, images, review_reasons
	if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
		images.append(image_asset(shape, assets_dir, djot_root, known_images))
		return titles, text_blocks, images, review_reasons
	if getattr(shape, "has_table", False):
		for row in shape.table.rows:
			cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
			if cells:
				line = " | ".join(djot_text(pptx_common.markdown_text(cell)) for cell in cells)
				text_blocks.append(pptx_common.TextBlock(shape.left, shape.top, ((0, line),), False))
		return titles, text_blocks, images, review_reasons
	if getattr(shape, "has_text_frame", False):
		line_values: list[tuple[int, str]] = []
		for paragraph in shape.text_frame.paragraphs:
			text = djot_text(pptx_common.paragraph_markdown(paragraph))
			if text:
				line_values.append((paragraph.level, text))
		lines = tuple(line_values)
		if not lines:
			if shape.shape_type != MSO_SHAPE_TYPE.PLACEHOLDER:
				review_reasons.append(f"ignored non-content shape type {shape.shape_type}")
			return titles, text_blocks, images, review_reasons
		if shape.shape_id == title_id:
			titles.extend(text for _level, text in lines)
		else:
			text_blocks.append(
				pptx_common.TextBlock(
					shape.left,
					shape.top,
					lines,
					pptx_common.is_subtitle_shape(shape),
				)
			)
		return titles, text_blocks, images, review_reasons
	if shape.shape_type != MSO_SHAPE_TYPE.PLACEHOLDER:
		review_reasons.append(f"ignored non-content shape type {shape.shape_type}")
	return titles, text_blocks, images, review_reasons


#============================================
def slide_notes(slide: object) -> tuple[str, ...]:
	"""Read source notes safely, but never render them as slide source."""
	has_notes_part = any(rel.reltype.endswith("/notesSlide") for rel in slide.part.rels.values())
	if not has_notes_part:
		return ()
	notes_frame = slide.notes_slide.notes_text_frame
	if notes_frame is None:
		return ()
	return tuple(line.strip() for line in notes_frame.text.splitlines() if line.strip())


#============================================
def extract_slides(
	presentation: object,
	assets_dir: pathlib.Path,
	djot_root: pathlib.PurePosixPath,
	expected_hidden: set[int] | None,
) -> tuple[list[pptx_common.SlideData], int]:
	"""Extract slides while preserving source order and authoritative visibility."""
	known_images: dict[str, str] = {}
	slides: list[pptx_common.SlideData] = []
	for source_index, slide in enumerate(presentation.slides, start=1):
		pptx_hidden = slide.element.get("show") == "0"
		hidden = pptx_hidden if expected_hidden is None else source_index in expected_hidden
		if expected_hidden is not None and pptx_hidden != hidden:
			raise RuntimeError(f"ODP and PPTX visibility disagree on slide {source_index}")
		if hidden:
			slides.append(pptx_common.SlideData(source_index, True, (), (), (), (), ()))
			continue
		titles: list[str] = []
		text_blocks: list[pptx_common.TextBlock] = []
		images: list[pptx_common.ImageAsset] = []
		review_reasons: list[str] = []
		title_id = pptx_common.title_shape_id(slide)
		for shape in slide.shapes:
			parts = shape_inventory(shape, title_id, assets_dir, djot_root, known_images)
			titles.extend(parts[0])
			text_blocks.extend(parts[1])
			images.extend(parts[2])
			review_reasons.extend(parts[3])
		line_count = sum(len(block.lines) for block in text_blocks)
		character_count = sum(
			len(text) for block in text_blocks for _level, text in block.lines
		)
		if line_count > 10 or character_count > 1200:
			review_reasons.append("dense text requires post-conversion polish")
		text_blocks.sort(key=lambda block: (block.top, block.left))
		images.sort(key=lambda image: (image.top, image.left))
		slides.append(
			pptx_common.SlideData(
				source_index=source_index,
				hidden=False,
				title_lines=tuple(titles),
				text_blocks=tuple(text_blocks),
				images=tuple(images),
				notes=slide_notes(slide),
				review_reasons=tuple(sorted(set(review_reasons))),
			)
		)
	return slides, len(known_images)


#============================================
def body_lines(slide: pptx_common.SlideData) -> list[str]:
	"""Render positioned text blocks as Djot lists and subtitles."""
	lines: list[str] = []
	for block in slide.text_blocks:
		for level, text in block.lines:
			if block.is_subtitle:
				lines.append(f"## {text}")
			else:
				lines.append(f"{'  ' * level}- {text}")
	return lines


#============================================
def image_djot(image: pptx_common.ImageAsset) -> str:
	"""Render one native Djot image, reserving Marp import syntax exactly."""
	return f"![{image.alt_text}]({image.markdown_path})"


#============================================
def titled_slide(slide: pptx_common.SlideData) -> tuple[list[str], list[str]]:
	"""Return source title lines and remaining body content."""
	texts = body_lines(slide)
	title_lines = list(slide.title_lines)
	if not title_lines and texts:
		title_lines = [texts.pop(0).removeprefix("- ")]
	if not title_lines:
		title_lines = [f"Slide {slide.source_index}"]
	return [f"# {' '.join(title_lines)}"], texts


#============================================
def render_slide(
	slide: pptx_common.SlideData,
	slide_width: int,
	is_first: bool,
) -> tuple[list[str], str]:
	"""Map one semantic slide to experimental layout and named Djot slots."""
	heading_lines, texts = titled_slide(slide)
	images = list(slide.images)
	if not images:
		if is_first and (not texts or all(text.startswith("## ") for text in texts)):
			return ["=== layout: title-slide", "", *heading_lines, "", *texts], "title-slide"
		if not texts:
			return ["=== layout: title-only", "", *heading_lines], "title-only"
		return ["=== layout: title-content", "", *heading_lines, "", "@body", "", *texts], "title-content"
	if len(images) == 1 and not texts:
		return [
			"=== layout: title-content", "", *heading_lines, "", "@body", "",
			image_djot(images[0]),
		], "title-content"
	if len(images) == 1:
		image = images[0]
		center = (image.left + image.width / 2) / slide_width
		if center >= 0.55:
			return [
				"=== layout: two-panels", "", *heading_lines, "", "@left", "", *texts,
				"", "@right", "", image_djot(image),
			], "two-panels"
		if center <= 0.45:
			return [
				"=== layout: two-panels", "", *heading_lines, "", "@left", "",
				image_djot(image), "", "@right", "", *texts,
			], "two-panels"
	image_lines = [image_djot(image) for image in images]
	if not texts:
		return ["=== layout: gallery", "", *heading_lines, "", "@gallery", "", *image_lines], "gallery"
	return [
		"=== layout: two-panels", "", *heading_lines, "", "@left", "", *texts,
		"", "@right", "", *image_lines,
	], "two-panels"


#============================================
def render_djot(
	slides: list[pptx_common.SlideData],
	slide_width: int,
) -> tuple[str, list[dict[str, object]]]:
	"""Render visible slide source and an auditable per-slide import report."""
	visible_slides = [slide for slide in slides if not slide.hidden]
	djot_lines: list[str] = []
	records: list[dict[str, object]] = []
	for visible_index, slide in enumerate(visible_slides):
		if visible_index:
			djot_lines.append("")
		slide_lines, layout = render_slide(slide, slide_width, visible_index == 0)
		djot_lines.extend(slide_lines)
		records.append(
			{
				"source_slide": slide.source_index,
				"layout": layout,
				"text_blocks": len(slide.text_blocks),
				"images": len(slide.images),
				"notes_omitted": len(slide.notes),
				"review_reasons": list(slide.review_reasons),
			}
		)
	return "\n".join(djot_lines).rstrip() + "\n", records


#============================================
def validate_output_path(output_path: pathlib.Path) -> None:
	"""Protect an established Djot destination from replacement."""
	if output_path.suffix.lower() != ".djot":
		raise ValueError("output must use the .djot extension")
	if output_path.exists():
		raise FileExistsError("output Djot source already exists; import will not overwrite it")
	asset_path = output_path.parent / "assets" / output_path.stem
	if asset_path.exists():
		raise FileExistsError("output asset directory already exists")


#============================================
def convert_pptx(
	input_path: pathlib.Path,
	output_path: pathlib.Path,
	*,
	expected_slide_count: int | None = None,
	expected_hidden: set[int] | None = None,
	source_name: str | None = None,
) -> ConversionSummary:
	"""Convert one trusted PPTX into new experimental extended-Djot source."""
	input_path = input_path.resolve()
	output_path = output_path.resolve()
	pptx_common.validate_pptx(input_path)
	validate_output_path(output_path)
	presentation = Presentation(input_path)
	if expected_slide_count is not None and len(presentation.slides) != expected_slide_count:
		raise RuntimeError("ODP and normalized PPTX slide counts disagree")
	output_path.parent.mkdir(parents=True, exist_ok=True)
	with tempfile.TemporaryDirectory(prefix=".pptx_to_djot_", dir=output_path.parent) as temporary_name:
		temporary_root = pathlib.Path(temporary_name)
		staging_assets = temporary_root / "assets"
		staging_assets.mkdir()
		djot_root = pathlib.PurePosixPath("assets") / output_path.stem
		slides, image_count = extract_slides(presentation, staging_assets, djot_root, expected_hidden)
		visible_slides = [slide for slide in slides if not slide.hidden]
		if not visible_slides:
			raise ValueError("presentation contains no visible slides")
		djot, records = render_djot(slides, presentation.slide_width)
		report = {
			"source": source_name or input_path.name,
			"slide_count": len(slides),
			"visible_slides": len(visible_slides),
			"hidden_slides": [slide.source_index for slide in slides if slide.hidden],
			"slides": records,
		}
		report_path = staging_assets / "import_report.json"
		report_path.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
		staging_djot = temporary_root / output_path.name
		staging_djot.write_text(djot, encoding="utf-8")
		final_assets = output_path.parent / "assets" / output_path.stem
		final_assets.parent.mkdir(exist_ok=True)
		os.replace(staging_assets, final_assets)
		os.replace(staging_djot, output_path)
	final_report = output_path.parent / "assets" / output_path.stem / "import_report.json"
	review_count = sum(bool(slide.review_reasons) for slide in slides if not slide.hidden)
	return ConversionSummary(
		visible_slides=len(visible_slides),
		editable_slides=len(visible_slides),
		hidden_slides=len(slides) - len(visible_slides),
		extracted_images=image_count,
		review_slides=review_count,
		output_path=output_path,
		report_path=final_report,
	)


#============================================
def parse_args() -> argparse.Namespace:
	"""Parse standalone PPTX-to-Djot importer arguments."""
	parser = argparse.ArgumentParser(description=__doc__)
	parser.add_argument("input_file", type=pathlib.Path, help="trusted source PPTX")
	parser.add_argument("--output", type=pathlib.Path, help="new extended-Djot source path")
	return parser.parse_args()


#============================================
def main() -> None:
	"""Run one standalone PPTX-to-Djot conversion."""
	args = parse_args()
	output_path = args.output or args.input_file.with_suffix(".djot")
	summary = convert_pptx(args.input_file, output_path)
	print(
		f"Converted {summary.visible_slides} visible slides: "
		f"{summary.editable_slides} editable, {summary.review_slides} layout review, "
		f"{summary.hidden_slides} hidden, {summary.extracted_images} content images"
	)
	print(f"Djot: {summary.output_path}")
	print(f"Import report: {summary.report_path}")


if __name__ == "__main__":
	main()
