"""Compact behavioral coverage for exceptional legacy visual planning."""

# Standard Library
import dataclasses
import types

# PIP3 modules
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.util import Inches

# Local modules
import pytest
import marp_lib.importers.legacy_slide_plan as plan
import marp_lib.importers.legacy_geometry as geometry
import marp_lib.importers.legacy_djot_emitter as emitter
import marp_lib.importers.legacy_new_visual_relations as visual_relations
import marp_lib.importers.legacy_topology as topology
import marp_lib.importers.pptx_to_djot as importer
import marp_lib.importers.pptx_to_marp as source


def bounds(left: float, top: float, right: float, bottom: float) -> plan.NormalizedBounds:
	"""Build one normalized rectangle for inline planner evidence."""
	return plan.NormalizedBounds(left, top, right, bottom)


def text(left: float, top: float, right: float, bottom: float, *, confidence: float = 0.0,
	kind: str = "auto-shape", ordinal: int = 1, paragraphs: int = 1,
	rotation: float = 0.0, fill: bool = False, line: bool = False,
	role: str | None = None, z_order: tuple[int, ...] = ()) -> plan.SourceTextRegion:
	"""Build one short source text shape without a fixture file."""
	return plan.SourceTextRegion(tuple((0, "Label") for _ in range(paragraphs)),
		bounds(left, top, right, bottom), False, confidence, False, kind, ordinal,
		rotation_degrees=rotation, has_positive_fill=fill, has_positive_line=line,
		placeholder_role=role, z_order=z_order)


def image(left: float, top: float, right: float, bottom: float, *, kind: str = "picture",
	ordinal: int = 1, z_order: tuple[int, ...] = ()) -> plan.SourceImageRegion:
	"""Build one source visual with stable provenance."""
	return plan.SourceImageRegion(f"source-{ordinal}", bounds(left, top, right, bottom), kind, ordinal, z_order=z_order)


def test_source_text_inventory_preserves_canonical_shape_rotation() -> None:
	"""Rotation remains normalized source evidence without affecting text extraction."""
	presentation = Presentation()
	slide = presentation.slides.add_slide(presentation.slide_layouts[5])
	shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
	shape.text = "Rotated label"
	shape.rotation = 315
	regions = importer.source_text_regions(slide, presentation.slide_width, presentation.slide_height)

	assert len(regions) == 1
	assert regions[0].rotation_degrees == -45.0


def test_source_visual_regions_preserve_top_level_and_group_z_paths() -> None:
	"""Visual source paths retain both top-level and grouped stack positions."""
	picture = lambda shape_id: types.SimpleNamespace(shape_type=MSO_SHAPE_TYPE.PICTURE, shape_id=shape_id,
		left=0, top=0, width=10, height=10)
	slide = types.SimpleNamespace(shapes=(picture(1), types.SimpleNamespace(
		shape_type=MSO_SHAPE_TYPE.GROUP, shapes=(picture(2),),
	)))

	regions = importer.source_visual_regions(slide, 100, 100)

	assert tuple(region.z_order for region in regions) == ((0,), (1, 0))


def test_degenerate_connectors_preserve_endpoints_and_planning_footprints() -> None:
	"""Horizontal and vertical source lines inflate only their zero planning axis."""
	vertical = geometry.degenerate_connector_footprint(500, 100, 0, 600, 20, 1000, 1000)
	horizontal = geometry.degenerate_connector_footprint(-100, 0, 600, 0, 1, 1000, 1000)

	assert vertical is not None and vertical.endpoints == geometry.NormalizedLineEndpoints(.5, .1, .5, .7)
	assert vertical.footprint == bounds(.495, .1, .505, .7)
	assert horizontal is not None and horizontal.endpoints == geometry.NormalizedLineEndpoints(0.0, 0.0, .5, 0.0)
	assert horizontal.footprint == bounds(0.0, 0.0, .5, .0005)


def test_source_inventory_retains_only_visible_degenerate_connector_metadata() -> None:
	"""The PPTX boundary supplies a planning footprint without changing source endpoints."""
	presentation = Presentation()
	slide = presentation.slides.add_slide(presentation.slide_layouts[5])
	connector = slide.shapes.add_connector(
		MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(2), Inches(5), Inches(2),
	)
	connector.line.color.rgb = RGBColor(0, 0, 0)
	connector.line.width = Inches(.02)
	slide.shapes.add_connector(
		MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(3), Inches(5), Inches(3),
	)
	regions = importer.source_visual_regions(slide, presentation.slide_width, presentation.slide_height)

	assert len(regions) == 1 and regions[0].source_kind == "connector"
	assert regions[0].source_line is not None
	assert regions[0].bounds.height > 0 and regions[0].source_line.endpoints.top == regions[0].source_line.endpoints.bottom


@pytest.mark.parametrize("case", ("unstroked", "nonconnector", "text", "point"))
def test_source_inventory_rejects_degenerate_connector_boundary_cases(case: str) -> None:
	"""Only a blank, visible stroked line with one zero axis enters visual inventory."""
	stroke = case in {"text", "point"}
	xml = '<a:prstGeom prst="line">' if case != "nonconnector" else "<p:sp>"
	line = types.SimpleNamespace(
		find=lambda name: object() if stroke and name.endswith("solidFill") else None,
		get=lambda name: "12700" if name == "w" else None,
	)
	properties = types.SimpleNamespace(
		find=lambda name: line if name.endswith("ln") else (object() if case == "nonconnector" else None),
	)
	shape = types.SimpleNamespace(
		shape_type=MSO_SHAPE_TYPE.AUTO_SHAPE, is_placeholder=False, has_table=False,
		has_text_frame=case == "text", text_frame=types.SimpleNamespace(
			paragraphs=(types.SimpleNamespace(text="Visible label"),) if case == "text" else (),
		), element=types.SimpleNamespace(
			xml=xml, find=lambda name: properties if name.endswith("spPr") else None,
		), shape_id=61, left=10, top=10, width=0, height=0 if case == "point" else 40,
	)

	assert not importer.source_visual_inventory(shape, 100, 100)


@pytest.mark.parametrize("geometry_values", (
	(0, 0, 0, 0, 10), (0, 0, 1, 1, 10), (0, 0, -1, 0, 10),
	(0, 0, 0, 1, 0), (0, 0, 0, 1, -1),
))
def test_degenerate_connector_normalization_rejects_invalid_geometry(
	geometry_values: tuple[int, int, int, int, int],
) -> None:
	"""Only one-axis degenerate, positively stroked connectors receive footprints."""
	line = geometry.degenerate_connector_footprint(*geometry_values, 1000, 1000)

	assert line is None


def test_degenerate_connectors_join_existing_vector_scaffold_without_consuming_prose() -> None:
	"""Inflated connector footprints participate in the existing scaffold relation."""
	first = geometry.degenerate_connector_footprint(300, 300, 400, 0, 10, 1000, 1000)
	second = geometry.degenerate_connector_footprint(300, 500, 400, 0, 10, 1000, 1000)
	assert first is not None and second is not None
	connectors = (
		plan.SourceImageRegion("source-connector-71", first.footprint, "connector", 71, first),
		plan.SourceImageRegion("source-connector-72", second.footprint, "connector", 72, second),
	)
	labels = (text(.30, .34, .38, .40, ordinal=73), text(.62, .34, .70, .40, ordinal=74),
		text(.46, .43, .54, .50, ordinal=75))
	exterior = text(.02, .30, .20, .65, confidence=1.0, kind="text", ordinal=76)
	slide_plan = plan.plan_slide((*labels, exterior), connectors)

	assert slide_plan.content_region is not None and slide_plan.content_region.kind == "vector-scaffold"
	assert slide_plan.content_region.image_regions == connectors
	assert connectors[0].source_line is not None and connectors[0].source_line.endpoints == first.endpoints
	assert exterior in slide_plan.slots[0].text_regions


def test_visual_exception_rules_preserve_only_proven_components() -> None:
	"""Tiny decoration, shared captions, coarse labels, and connectors stay explicit."""
	tiny = image(.02, .02, .03, .03, kind="vector", ordinal=9)
	long = image(.02, .02, .03, .20, kind="vector", ordinal=10)
	assert plan.plan_slide((), (tiny,)).omitted_vectors == (tiny,)
	assert plan.plan_slide((), (long,)).review_vectors == (long,)
	first, second = image(.10, .30, .35, .55, ordinal=1), image(.45, .31, .70, .56, ordinal=2)
	third = image(.85, .29, .95, .54, ordinal=13)
	caption = text(.18, .58, .88, .64, ordinal=3)
	row = plan.plan_slide((caption,), (first, second, third)).content_region
	assert row is not None and row.kind == "shared-figure-row" and row.image_regions == (first, second, third)
	picture = image(.10, .20, .80, .80, ordinal=4)
	labels = (text(.20, .35, .28, .42, ordinal=5), text(.65, .60, .73, .67, ordinal=6))
	connector = image(.15, .82, .75, .83, kind="connector", ordinal=7)
	diagram = plan.plan_slide(labels, (picture, connector)).content_region
	assert diagram is not None and connector in diagram.image_regions and diagram.bounds.bottom == connector.bounds.bottom
	coarse = text(.19, .34, .56, .61, confidence=1.0, kind="text", ordinal=8, paragraphs=2)
	mixed = plan.plan_slide((coarse, text(.25, .40, .35, .48, ordinal=11)), (image(.20, .35, .55, .60, ordinal=12),)).content_region
	assert mixed is not None and coarse in mixed.text_regions and mixed.bounds == coarse.bounds


def test_repeated_figures_vector_labels_and_flat_below_answers_stay_structural() -> None:
	"""New relation classifiers consume only their complete symmetric source evidence."""
	left, right = image(.15, .30, .40, .55, ordinal=20), image(.60, .30, .85, .55, ordinal=21)
	outer_labels = (text(.05, .37, .16, .45, ordinal=22), text(.84, .37, .95, .45, ordinal=23))
	repeated = plan.plan_slide(outer_labels, (left, right)).content_region
	assert repeated is not None and repeated.kind == "repeated-labeled-figure"
	first, second = image(.20, .30, .55, .60, kind="vector", ordinal=24), image(.50, .30, .80, .60, kind="vector", ordinal=25)
	labels = (text(.23, .34, .30, .40, ordinal=26), text(.70, .50, .77, .56, ordinal=27))
	contained = text(.35, .42, .45, .52, confidence=1.0, ordinal=28, paragraphs=3)
	unrelated = text(.02, .75, .15, .82, confidence=1.0, ordinal=29)
	full_vector_plan = plan.plan_slide((*labels, contained, unrelated), (first, second))
	vector_plan = full_vector_plan.content_region
	assert vector_plan is not None and contained in vector_plan.text_regions and unrelated in full_vector_plan.slots[0].text_regions
	question = plan.SourceTextRegion(((0, "Prompt"), (1, "Choice")), bounds(.08, .20, .82, .68), False, 1.0)
	answer = plan.SourceTextRegion(((0, "Answer"), (0, "Why")), bounds(.60, .70, .80, .76), False, 0.0)
	assert plan.multiple_choice_plan((question, answer), ()) is not None


def test_local_figure_heading_stays_editable_above_an_unchanged_crop() -> None:
	"""One short exterior auto-shape becomes local H2 before its source-region asset."""
	heading = text(.25, .13, .35, .17, ordinal=31)
	labels = (text(.30, .40, .37, .47, ordinal=32), text(.65, .60, .72, .67, ordinal=33))
	picture = image(.20, .20, .80, .80, ordinal=34)
	slide_plan = plan.plan_slide((heading, *labels), (picture,))
	content = slide_plan.content_region
	planned = importer.PlannedSlide(source.SlideData(1, False, (), (), (), (), ()), slide_plan)
	lines, layout, _reasons = emitter.render_planned_slide(planned, {(1, content.asset_key): "assets/crop.png"}, False)

	assert content is not None and content.local_heading is heading and content.bounds == picture.bounds
	assert layout == "one-panel" and lines[-3:] == ["## Label", "", "![Coupled source region](assets/crop.png)"]


@pytest.mark.parametrize("case", (
	"inside", "title", "protected", "side", "below", "multi", "prose", "text", "image", "table",
))
def test_local_figure_heading_rejects_competing_or_nonheading_evidence(case: str) -> None:
	"""One bounded matrix keeps the heading relation limited to its geometric contract."""
	picture = image(.20, .20, .80, .80, ordinal=40)
	heading = text(.30, .13, .40, .17, ordinal=41, paragraphs=2 if case == "multi" else 1,
		kind="text" if case == "prose" else "auto-shape")
	content = plan.ContentRegionPlan(
		"crop", picture.bounds, (), (picture,),
		protected_text_shape_ids=(heading.source_ordinal,) if case == "protected" else (),
	)
	if case == "inside":
		heading = text(.30, .30, .40, .35, ordinal=41)
	if case == "side":
		heading = text(.02, .13, .12, .17, ordinal=41)
	if case == "below":
		heading = text(.30, .82, .40, .86, ordinal=41)
	if case == "text":
		extras = (text(.55, .13, .65, .17, ordinal=42),)
	else:
		extras = ()
	title = plan.TitleDecision(heading if case == "title" else None, "test")
	images = (picture, image(.82, .20, .92, .30, ordinal=43)) if case == "image" else (picture,)
	tables = (object(),) if case == "table" else ()

	assert plan.local_figure_heading(content, (heading, *extras), images, title, tables) is None


def test_rotated_vector_labels_keep_only_one_complete_connected_figure() -> None:
	"""A rotated three-label lane retains its vector scaffold and exterior prose stays native."""
	vectors = (image(.30, .38, .70, .42, kind="vector", ordinal=51),
		image(.30, .53, .70, .57, kind="vector", ordinal=52))
	labels = (text(.30, .30, .70, .366583, ordinal=53, rotation=18),
		text(.30, .4567, .70, .50, ordinal=54), text(.30, .60, .70, .65, ordinal=55))
	exterior = text(.02, .30, .20, .60, confidence=1.0, kind="text", ordinal=56)
	slide_plan = plan.plan_slide((*labels, exterior), vectors)
	content = slide_plan.content_region

	assert content is not None and content.kind == "rotated-vector-label"
	assert content.text_regions == labels and content.image_regions == vectors
	assert exterior in slide_plan.slots[0].text_regions
	assert not any(region in labels for slot in slide_plan.slots for region in slot.text_regions)
	assert not any(image_region in vectors for slot in slide_plan.slots for image_region in slot.image_regions)


@pytest.mark.parametrize(
	"case", (
		"unrotated", "single-vector", "single-label", "disconnected", "title", "identity",
		"table", "full-slide", "competing",
	),
)
def test_rotated_vector_labels_require_complete_non_title_evidence(case: str) -> None:
	"""Rotation is provenance, not permission to crop incomplete vector arrangements."""
	vectors = (image(.30, .38, .70, .42, kind="vector", ordinal=61),
		image(.30, .53, .70, .57, kind="vector", ordinal=62))
	labels = (text(.30, .30, .70, .366, ordinal=63, rotation=16),
		text(.30, .456, .70, .50, ordinal=64))
	if case == "unrotated":
		labels = tuple(dataclasses.replace(label, rotation_degrees=0.0) for label in labels)
	if case == "single-vector":
		vectors = vectors[:1]
	if case == "single-label":
		labels = labels[:1]
	if case == "disconnected":
		labels = (labels[0], dataclasses.replace(labels[1], bounds=bounds(.80, .80, .90, .86)))
	if case == "title":
		labels = (
			dataclasses.replace(labels[0], bounds=bounds(.30, .04, .70, .12), title_identity=True,
				placeholder_confidence=1.0), labels[1],
		)
	if case == "identity":
		labels = (
			dataclasses.replace(labels[0], bounds=bounds(.30, .04, .70, .12), title_identity=True), labels[1],
		)
	if case == "table":
		labels = (labels[0], dataclasses.replace(labels[1], source_kind="table"))
	if case == "full-slide":
		vectors = (
			dataclasses.replace(vectors[0], bounds=bounds(0.0, .38, .70, .42)),
			dataclasses.replace(vectors[1], bounds=bounds(.30, .53, 1.0, .57)),
		)
		labels = (
			dataclasses.replace(labels[0], bounds=bounds(0.0, 0.0, .70, .366)),
			dataclasses.replace(labels[1], bounds=bounds(.30, .576, 1.0, 1.0)),
		)
	if case == "competing":
		labels = (*labels, text(.45, .44, .55, .48, confidence=1.0, kind="text", ordinal=65))
	content = plan.plan_slide(labels, vectors).content_region

	assert content is None or content.kind != "rotated-vector-label"


def test_styled_callout_requires_connector_and_matching_left_edge() -> None:
	"""A connector crop grows only for one matching, explicitly styled editable callout."""
	vector = image(.20, .25, .70, .70, kind="connector", ordinal=80)
	inside = text(.30, .35, .45, .42, ordinal=81, paragraphs=4, fill=True, line=True)
	callout = text(.31, .72, .46, .79, ordinal=82, fill=True, line=True)
	content = plan.ContentRegionPlan("visual", vector.bounds, (inside,), (vector,))
	relation = visual_relations.styled_callout_extension(
		content, (inside, callout), plan.TitleDecision(None, "test"),
	)

	assert relation is not None and relation.text_regions == (inside, callout)
	assert visual_relations.styled_callout_extension(
		content, (inside, dataclasses.replace(callout, has_positive_line=False)),
		plan.TitleDecision(None, "test"),
	) is None
	assert visual_relations.styled_callout_extension(
		dataclasses.replace(content, text_regions=(dataclasses.replace(
			inside, paragraphs=((0, "Label"),) * 5,
		),)), (inside, callout),
		plan.TitleDecision(None, "test"),
	) is None
	assert visual_relations.styled_callout_extension(
		content, (inside, dataclasses.replace(callout, paragraphs=((0, "Label"),) * 3)),
		plan.TitleDecision(None, "test"),
	) is None
	assert visual_relations.styled_callout_extension(
		content, (inside, dataclasses.replace(callout, bounds=bounds(.38, .72, .53, .79))),
		plan.TitleDecision(None, "test"),
	) is None
	assert not visual_relations.styled_callout_candidate(
		text(.31, .60, .46, .68, ordinal=83, fill=True, line=True), content.bounds,
	)
	protected_title = text(.30, .35, .45, .42, confidence=1.0, ordinal=84)
	protected_content = dataclasses.replace(content, protected_text_shape_ids=(84,))
	protected_relation = visual_relations.styled_callout_extension(
		protected_content, (protected_title, inside, callout),
		plan.TitleDecision(protected_title, "test"),
	)

	assert protected_relation is not None and protected_relation.protected_text_shape_ids == (84,)
	candidate_title = text(.32, .73, .44, .78, confidence=1.0, ordinal=85)
	candidate_protected = dataclasses.replace(content, protected_text_shape_ids=(85,))
	assert visual_relations.styled_callout_extension(
		candidate_protected, (candidate_title, inside, callout),
		plan.TitleDecision(candidate_title, "test"),
	) is None


def test_coarse_body_styled_key_plans_native_body_and_two_panel_asset() -> None:
	"""A transparent object placeholder retains a styled inset key as a crop asset."""
	body = text(.10, .10, .90, .90, confidence=1.0, ordinal=90, role="OBJECT", z_order=(1,))
	key = text(.68, .11, .82, .19, ordinal=91, fill=True, line=True, z_order=(2,))
	title = text(.10, .02, .90, .08, confidence=1.0, ordinal=89, role="TITLE", z_order=(0,))
	slide_plan = plan.plan_slide((title, body, key), ())
	content = slide_plan.content_region
	planned = importer.PlannedSlide(source.SlideData(1, False, (), (), (), (), ()), slide_plan)
	lines, layout, reasons = emitter.render_planned_slide(
		planned, {(1, content.asset_key): "assets/key.png"}, False,
	)

	assert content is not None and content.kind == "styled-inset-key"
	assert slide_plan.title.region is title
	assert slide_plan.slots[0].text_regions == (body,)
	assert layout == "two-panels" and "@left" in lines and "@right" in lines and not reasons


def test_coarse_body_picture_inset_keeps_both_members_native() -> None:
	"""A small late picture inset occupies its side while the transparent body stays editable."""
	title = text(.10, .02, .90, .08, confidence=1.0, ordinal=89, role="TITLE", z_order=(0,))
	body = text(.10, .10, .90, .90, confidence=1.0, ordinal=90, role="OBJECT", z_order=(1,))
	picture = image(.10, .30, .25, .60, ordinal=91, z_order=(2,))
	slide_plan = plan.plan_slide((title, body), (picture,))
	data = source.SlideData(1, False, (), (), (source.ImageAsset(0, 0, 1, 1, "source-91", "Inset"),), (), ())
	lines, layout, reasons = emitter.render_planned_slide(importer.PlannedSlide(data, slide_plan), {}, False)

	assert slide_plan.title.region is title and slide_plan.content_region is None
	assert tuple(slot.name for slot in slide_plan.slots) == ("left", "right")
	assert layout == "two-panels" and "@left" in lines and "@right" in lines and not reasons


def test_caption_unit_and_footer_use_existing_two_plus_one_tolerance() -> None:
	"""One body, paired figure, and shallow overlapping footer remain native."""
	title = text(.05, .02, .95, .08, confidence=1.0, ordinal=80, role="TITLE")
	body = text(.68, .20, .92, .70, ordinal=81)
	picture = image(.10, .20, .42, .52, ordinal=82)
	caption = text(.13, .511, .39, .56, ordinal=83)
	footer = text(.10, .54, .68, .61, kind="text-box", ordinal=84)
	slide_plan = plan.plan_slide((title, body, caption, footer), (picture,))
	data = source.SlideData(1, False, (), (), (source.ImageAsset(0, 0, 1, 1, "source-82", "Figure"),), (), ())
	planned = importer.PlannedSlide(data, slide_plan)
	lines, layout, reasons = emitter.render_planned_slide(planned, {}, False)
	components, _reasons = emitter.emit_components(planned, None)
	permissions = emitter.overlap_permissions(components)

	assert layout == "two-plus-one-panels" and not reasons
	assert len(components) == 3 and any(len(item.member_footprints) == 2 for item in components)
	assert [item.relation for item in permissions] == ["caption-footer-padding"] and "![Figure](source-82)" in lines


@pytest.mark.parametrize("case", ("opaque", "lined", "role", "extra", "large", "vertical", "overhang", "center", "z-order", "foreign-title"))
def test_coarse_body_picture_inset_rejects_nonunique_geometry(case: str) -> None:
	"""Only the sole transparent OBJECT body with one bounded side inset qualifies."""
	title = text(.10, .02, .90, .08, confidence=1.0, ordinal=89, role="TITLE")
	body = text(.10, .10, .90, .90, confidence=1.0, ordinal=90, role="OBJECT", z_order=(1,))
	picture = image(.10, .30, .25, .60, ordinal=91, z_order=(2,))
	if case == "opaque": body = dataclasses.replace(body, has_positive_fill=True)
	if case == "lined": body = dataclasses.replace(body, has_positive_line=True)
	if case == "role": body = dataclasses.replace(body, placeholder_role="BODY")
	if case == "large": picture = image(.10, .20, .70, .80, ordinal=91)
	if case == "vertical": picture = image(.10, .05, .25, .60, ordinal=91)
	if case == "overhang": picture = image(.05, .30, .25, .60, ordinal=91)
	if case == "center": picture = image(.40, .30, .60, .60, ordinal=91)
	if case == "z-order": picture = dataclasses.replace(picture, z_order=(0,))
	images = (picture,) if case != "extra" else (picture, image(.30, .30, .35, .35, ordinal=92))
	selected = title if case != "foreign-title" else text(.10, .02, .90, .08, confidence=1.0, ordinal=99, role="TITLE")

	assert visual_relations.coarse_body_picture_inset((title, body), images, plan.TitleDecision(selected, "test")) is None


def test_dominant_picture_narrative_uses_the_complete_visual_membership() -> None:
	"""A large picture, upper-right inset, and lower prose form one bounded crop."""
	dominant = image(.10, .12, .80, .72, ordinal=100)
	inset = image(.65, .16, .75, .26, ordinal=101)
	narrative = text(.12, .74, .78, .82, ordinal=102)
	content = plan.plan_slide((narrative,), (dominant, inset)).content_region

	assert content is not None and content.kind == "dominant-image-narrative"
	assert content.text_regions == (narrative,) and content.image_regions == (dominant, inset)


def test_dominant_picture_narrative_protects_title_overlapped_only_by_union() -> None:
	"""A full-width crop protects a title when only its union AABB reaches the title."""
	dominant = image(.10, .12, .80, .72, ordinal=103)
	inset = image(.65, .09, .75, .11, ordinal=104)
	narrative = text(.12, .74, .78, .82, ordinal=105)
	title_region = text(.78, .08, .95, .115, confidence=1.0, ordinal=106)
	title = plan.TitleDecision(title_region, "test")
	relation = visual_relations.dominant_image_narrative((title_region, narrative),
		(dominant, inset), title)

	assert relation is not None and relation.protected_text_shape_ids == (106,)


def test_visual_relations_reject_exact_full_crop_bounds() -> None:
	"""Exceptional relations never rasterize an exact full-slide region."""
	relation = visual_relations.VisualRelationMembers((), (), bounds(0.0, 0.0, 1.0, 1.0))

	assert not visual_relations.relation_is_safe(relation, plan.TitleDecision(None, "test"))


def test_repeated_figures_accept_mutual_labels_on_the_same_side() -> None:
	"""Both repeated figures may use their left-side labels when each pairing is unique."""
	left, right = image(.20, .30, .40, .55, ordinal=110), image(.60, .30, .80, .55, ordinal=111)
	labels = (text(.11, .37, .19, .45, ordinal=112), text(.51, .37, .59, .45, ordinal=113))
	content = plan.repeated_labeled_figure_region(labels, (left, right), plan.TitleDecision(None, "test"))

	assert content is not None and content.kind == "repeated-labeled-figure"
	assert content.text_regions == labels and content.image_regions == (left, right)


def test_visual_sequence_and_topology_viability_suppress_nonunique_crops() -> None:
	"""A qualified sequence is withheld when an existing native topology can represent it."""
	images = (image(.15, .15, .35, .30, ordinal=120), image(.30, .32, .50, .47, ordinal=121),
		image(.45, .49, .65, .64, ordinal=122))
	title = plan.TitleDecision(None, "test")
	relation = visual_relations.coupled_visual_sequence((), images, title)

	assert relation is not None and relation.image_regions == images
	assert visual_relations.coupled_visual_sequence((), images, title, lambda _relation: True) is None


def test_single_interior_overlay_label_uses_one_named_crop() -> None:
	"""One large source picture and its sole interior label remain spatially coupled."""
	label = text(.42, .42, .58, .50, confidence=1.0, role="BODY", ordinal=205)
	picture = image(.10, .20, .90, .80, ordinal=206)

	planned = plan.plan_slide((label,), (picture,))

	assert planned.content_region is not None
	assert planned.content_region.kind == "single-interior-overlay-label"
	assert planned.content_region.text_regions == (label,)
	assert planned.content_region.image_regions == (picture,)


def test_single_interior_overlay_label_protects_overlapping_title() -> None:
	"""A title overlapping only the picture remains outside its full protected crop."""
	title = plan.SourceTextRegion(((0, "Title"),), bounds(.10, .05, .90, .15), False, 1.0, True, "text", 204)
	label = text(.42, .42, .58, .50, ordinal=205)
	picture = image(.10, .10, .90, .80, ordinal=206)

	planned = plan.plan_slide((title, label), (picture,))

	assert planned.content_region is not None and planned.content_region.bounds == picture.bounds
	assert planned.content_region.protected_text_shape_ids == (204,)
	assert planned.title.region is title and title not in planned.content_region.text_regions


@pytest.mark.parametrize("labels,picture_bounds,title_overlap,extra_vector", (
	(((.01, .42, .09, .50),), (.10, .20, .90, .80), False, False),
	(((.05, .42, .15, .50),), (.10, .20, .90, .80), False, False),
	(((.10, .82, .90, .90),), (.10, .20, .90, .80), False, False),
	(((.42, .42, .58, .50), (.60, .42, .72, .50)), (.10, .20, .90, .80), False, False),
	(((.42, .42, .58, .50),), (.10, .20, .50, .60), False, False),
	(((.42, .42, .58, .50),), (0, 0, 1, 1), False, False),
	(((.42, .12, .58, .22),), (.10, .10, .90, .80), True, False),
	(((.42, .42, .58, .50),), (.10, .20, .90, .80), False, True),
))
def test_single_interior_overlay_label_rejects_nonunique_or_noninterior_geometry(
		labels: tuple[tuple[float, float, float, float], ...], picture_bounds: tuple[float, float, float, float], title_overlap: bool, extra_vector: bool,
) -> None:
	"""Adjacent, exterior, multiple, small, and full-slide candidates stay uncoupled."""
	regions = tuple(text(*item, ordinal=210 + index) for index, item in enumerate(labels))
	title = plan.SourceTextRegion(((0, "Title"),), bounds(.10, .05, .90, .15), False, 1.0, True, "text", 299)
	visuals = (image(*picture_bounds, ordinal=220),) + (() if not extra_vector else (image(.02, .02, .03, .03, kind="vector", ordinal=221),))
	planned = plan.plan_slide((title, *regions) if title_overlap else regions, visuals)

	assert planned.content_region is None or planned.content_region.kind != "single-interior-overlay-label"


def test_visual_sequence_rejects_vertical_overlap() -> None:
	"""A sequence requires each later picture to begin at or below the prior bottom."""
	images = (image(.10, .10, .30, .30, ordinal=123), image(.35, .29, .55, .45, ordinal=124),
		image(.60, .47, .80, .63, ordinal=125))

	assert visual_relations.coupled_visual_sequence((), images, plan.TitleDecision(None, "test")) is None


@pytest.mark.parametrize("source_bounds", (
	((.58, .56, .81, .68), (.49, .42, .69, .63)),
	((.10, .10, .30, .30), (.40, .40, .60, .60), (.70, .70, .90, .90)),
))
def test_shared_topology_agrees_with_live_registry_for_ambiguous_and_nonmatch_geometry(
		source_bounds: tuple[tuple[float, float, float, float], ...],
) -> None:
	"""Both shared topology entry points reject ambiguous or unrepresentable registry geometry."""
	regions = tuple(bounds(*item) for item in source_bounds)

	assert topology.ordinary_layout_match(regions) is None
	assert not topology.ordinary_layout_viable(regions)


def test_shared_topology_agrees_with_live_registry_for_unique_geometry() -> None:
	"""The matcher and viability use the same live two-panel registry policy."""
	regions = (bounds(.10, .20, .50, .80), bounds(.55, .20, .95, .80))

	assert topology.ordinary_layout_match(regions) == ("two-panels", (0, 1))
	assert topology.ordinary_layout_viable(regions)
