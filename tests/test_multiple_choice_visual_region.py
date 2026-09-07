"""Focused behavior for multiple-choice source-region overlays."""

# Standard Library
import pathlib

# PIP3 modules
import pytest
from pptx import Presentation
from pptx.util import Inches

# Local modules
import marp_lib.importers.legacy_djot_emitter as legacy_djot_emitter
import marp_lib.importers.legacy_slide_plan as legacy_slide_plan
import marp_lib.importers.pptx_to_djot as pptx_to_djot
import marp_lib.importers.pptx_to_marp as pptx_to_marp
import marp_lib.importers.source_region_render as source_region_render
import marp_lib.djot_parser


#============================================
def test_multiple_choice_overlay_crop_protects_native_question_and_answer() -> None:
	"""An associated question overlay becomes one crop before the editable prompt."""
	question = legacy_slide_plan.SourceTextRegion(
		((0, "Prompt"), (1, "Choice one"), (1, "Choice two")),
		legacy_slide_plan.NormalizedBounds(0.08, 0.20, 0.82, 0.68), False, 1.0,
		source_ordinal=4,
	)
	answer = legacy_slide_plan.SourceTextRegion(
		((0, "Answer"),), legacy_slide_plan.NormalizedBounds(0.60, 0.63, 0.80, 0.67),
		False, 0.0, source_ordinal=8,
	)
	overlay = legacy_slide_plan.SourceTextRegion(
		((0, "Overlay label"),), legacy_slide_plan.NormalizedBounds(0.35, 0.26, 0.55, 0.32),
		False, 0.0, False, "auto-shape", 12,
	)
	image = legacy_slide_plan.SourceImageRegion(
		"assets/figure.png", legacy_slide_plan.NormalizedBounds(0.30, 0.22, 0.70, 0.32),
		source_ordinal=15,
	)
	plan = legacy_slide_plan.plan_slide((question, answer, overlay), (image,))
	choice = plan.multiple_choice
	assert choice is not None and choice.question_visual_region is not None
	assert choice.question_visual_region.text_regions == (overlay,)
	assert choice.question_visual_region.protected_text_shape_ids == (4, 8)
	planned = pptx_to_djot.PlannedSlide(
		pptx_to_marp.SlideData(3, False, (), (), (), (), ()), plan, visible_page_index=2,
	)

	lines, layout, _reasons = legacy_djot_emitter.render_planned_slide(
		planned, {(3, "multiple-choice-visual-1"): "assets/deck/overlay.png"}, False,
	)

	assert layout == "multiple-choice"
	assert lines.index("![Question visual region](assets/deck/overlay.png)") < lines.index("Prompt")
	with pytest.raises(ValueError, match="source slide 3.*required renderer asset"):
		legacy_djot_emitter.render_planned_slide(planned, {}, False)
	unrelated = legacy_slide_plan.SourceTextRegion(
		((0, "Unrelated"),), legacy_slide_plan.NormalizedBounds(0.01, 0.01, 0.06, 0.06),
		False, 0.0, False, "auto-shape", 20,
	)
	assert legacy_slide_plan.plan_slide((question, answer, unrelated), (image,)).multiple_choice is None


#============================================
def test_protected_variant_clears_multiple_structural_text_shapes(tmp_path: pathlib.Path) -> None:
	"""One private crop clone can clear both editable multiple-choice structures."""
	source_path = tmp_path / "normalized.pptx"
	presentation = Presentation()
	slide = presentation.slides.add_slide(presentation.slide_layouts[6])
	question = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
	answer = slide.shapes.add_textbox(Inches(6), Inches(5), Inches(2), Inches(1))
	question.text = "Prompt and choices"
	answer.text = "Answer"
	presentation.save(source_path)
	request = source_region_render.SourceRegionRequest("multiple-choice-visual-1", 1, 1,
		source_region_render.NormalizedRegion(0.1, 0.1, 0.9, 0.9),
		(question.shape_id, answer.shape_id))
	clone = Presentation(source_region_render.prepare_protected_variant(source_path, (request,), tmp_path))
	shapes = {shape.shape_id: shape for shape in clone.slides[0].shapes}

	assert shapes[question.shape_id].text == ""
	assert shapes[answer.shape_id].text == ""


#============================================
@pytest.mark.parametrize("visuals", [
	(legacy_slide_plan.SourceImageRegion(
		"source-connector-9", legacy_slide_plan.NormalizedBounds(0.01, 0.01, 0.05, 0.05), "connector", 9,
	),),
	(
		legacy_slide_plan.SourceImageRegion(
			"assets/figure.png", legacy_slide_plan.NormalizedBounds(0.30, 0.22, 0.70, 0.32), "picture", 10,
		),
		legacy_slide_plan.SourceImageRegion(
			"source-connector-11", legacy_slide_plan.NormalizedBounds(0.01, 0.01, 0.05, 0.05), "connector", 11,
		),
	),
])
def test_multiple_choice_declines_unconsumed_connector_visuals(
	visuals: tuple[legacy_slide_plan.SourceImageRegion, ...],
) -> None:
	"""Connector visuals never disappear behind a recognized structural question."""
	question = legacy_slide_plan.SourceTextRegion(
		((0, "Prompt"), (1, "Choice one")),
		legacy_slide_plan.NormalizedBounds(0.08, 0.20, 0.82, 0.68), False, 1.0, source_ordinal=4,
	)
	answer = legacy_slide_plan.SourceTextRegion(
		((0, "Answer"),), legacy_slide_plan.NormalizedBounds(0.72, 0.78, 0.92, 0.86),
		False, 0.0, source_ordinal=8,
	)

	assert legacy_slide_plan.plan_slide((question, answer), visuals).multiple_choice is None


#============================================
def test_multiple_choice_two_paragraph_answer_survives_emission_and_parse(tmp_path: pathlib.Path) -> None:
	"""A flat two-paragraph answer remains two source paragraphs in authored order."""
	question = legacy_slide_plan.SourceTextRegion(
		((0, "Prompt"), (1, "Choice one")),
		legacy_slide_plan.NormalizedBounds(0.08, 0.20, 0.82, 0.68), False, 1.0, source_ordinal=4,
	)
	answer = legacy_slide_plan.SourceTextRegion(
		((0, "First answer paragraph"), (0, "Second answer paragraph")),
		legacy_slide_plan.NormalizedBounds(0.72, 0.78, 0.92, 0.88), False, 0.0, source_ordinal=8,
	)
	choice = legacy_slide_plan.MultipleChoicePlan(question, answer, None, "test")
	plan = legacy_slide_plan.LegacySlidePlan(
		legacy_slide_plan.TitleDecision(None, "multiple-choice"), (), multiple_choice=choice,
	)
	planned = pptx_to_djot.PlannedSlide(
		pptx_to_marp.SlideData(1, False, (), (), (), (), ()), plan, visible_page_index=1,
	)

	lines, _layout, _reasons = legacy_djot_emitter.render_planned_slide(planned, {}, False)
	source_path = tmp_path / "two-answer.djot"
	source_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
	deck = marp_lib.djot_parser.parse_deck(source_path)

	assert lines[-3:] == ["First answer paragraph", "", "Second answer paragraph"]
	assert len(deck.slides[0].cells[1].blocks) == 2
