"""Fast behavioral coverage for native editable PPTX reveal timing."""

# Standard Library
import pathlib
import zipfile
import xml.etree.ElementTree

# PIP3 Modules
import pytest
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement

# Local Modules
import marp_lib.native_export
import marp_lib.native_model
import marp_lib.pptx_animation


PRESENTATION_NAMESPACES = {
	"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
	"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


#============================================
def render_djot(tmp_path: pathlib.Path, source: str) -> xml.etree.ElementTree.Element:
	"""Render one inline Djot slide and return its saved native slide XML."""
	input_path = tmp_path / "animation.djot"
	output_path = tmp_path / "animation.pptx"
	input_path.write_text(source, encoding="utf-8")
	deck = marp_lib.native_export.parse_deck(input_path)
	marp_lib.native_export.render_native_pptx(deck, output_path)
	with zipfile.ZipFile(output_path) as archive:
		root = xml.etree.ElementTree.fromstring(archive.read("ppt/slides/slide1.xml"))
	return root


#============================================
def text_by_shape_id(root: xml.etree.ElementTree.Element) -> dict[str, str]:
	"""Map saved native shape IDs to their editable text for target checks."""
	result = {}
	for shape in root.findall(".//p:sp", PRESENTATION_NAMESPACES):
		shape_id = shape.find(".//p:cNvPr", PRESENTATION_NAMESPACES).attrib["id"]
		words = [text.text or "" for text in shape.findall(".//a:t", PRESENTATION_NAMESPACES)]
		result[shape_id] = "".join(words)
	return result


#============================================
def direct_reveal_deck(tmp_path: pathlib.Path) -> marp_lib.native_model.Deck:
	"""Build direct IR for backend-only fade coverage without expanding Djot grammar."""
	location = marp_lib.native_model.SourceLocation(tmp_path / "fade.ir", 1)
	appear = marp_lib.native_model.Reveal(marp_lib.native_model.RevealEffect.APPEAR,
		marp_lib.native_model.RevealSequence.OBJECT)
	fade = marp_lib.native_model.Reveal(marp_lib.native_model.RevealEffect.FADE,
		marp_lib.native_model.RevealSequence.OBJECT)
	body = marp_lib.native_model.Cell(location, (
		marp_lib.native_model.Paragraph(location, (marp_lib.native_model.Text("Appear"),), appear),
		marp_lib.native_model.Paragraph(location, (marp_lib.native_model.Text("Fade"),), fade),
	), "body")
	slide = marp_lib.native_model.Slide(location, "one-panel", None, False, (), (), (body,))
	return marp_lib.native_model.Deck(location.path, tmp_path, tmp_path, "Effects", False, (slide,), {})


#============================================
def test_native_timing_uses_one_root_and_source_ordered_shape_targets(tmp_path: pathlib.Path) -> None:
	"""Prefix and terminal object reveals become ordered click targets on one timing root."""
	root = render_djot(tmp_path, "=== layout: one-panel\n@body\n=> appear\nFirst\n\n- Last\n<= appear\n")
	timings = root.findall("p:timing", PRESENTATION_NAMESPACES)
	targets = [target.attrib["spid"] for target in root.findall(".//p:spTgt", PRESENTATION_NAMESPACES)]
	assert len(timings) == 1 and [text_by_shape_id(root)[target] for target in targets] == ["First", "Last"]


#============================================
def test_cascade_targets_top_level_list_items_and_descendants_in_one_outline(tmp_path: pathlib.Path) -> None:
	"""Top-level cascade clicks use zero-based inclusive ranges within one editable list."""
	root = render_djot(tmp_path, "=== layout: one-panel\n@body\n=> cascade appear\n- Parent\n\n  - Child\n- Second\n")
	targets = root.findall(".//p:spTgt", PRESENTATION_NAMESPACES)
	ranges = [(node.attrib["st"], node.attrib["end"]) for node in root.findall(".//p:pRg", PRESENTATION_NAMESPACES)]
	assert [target.attrib["spid"] for target in targets] == [targets[0].attrib["spid"]] * 2 and ranges == [("0", "1"), ("2", "2")]


#============================================
def test_appear_and_backend_fade_write_their_distinct_native_effects(tmp_path: pathlib.Path) -> None:
	"""Appear uses visibility while direct IR fade retains its native entrance effect."""
	output_path = tmp_path / "effects.pptx"
	marp_lib.native_export.render_native_pptx(direct_reveal_deck(tmp_path), output_path)
	with zipfile.ZipFile(output_path) as archive:
		root = xml.etree.ElementTree.fromstring(archive.read("ppt/slides/slide1.xml"))
	appear = root.findall(".//p:set/p:to/p:strVal", PRESENTATION_NAMESPACES)
	visibility = root.findall(".//p:set/p:cBhvr/p:attrNameLst/p:attrName", PRESENTATION_NAMESPACES)
	fade = root.findall(".//p:animEffect", PRESENTATION_NAMESPACES)
	assert ([node.attrib["val"] for node in appear] == ["visible"] and [node.text for node in visibility] == ["style.visibility"] and
		[(node.attrib["filter"], node.attrib["transition"]) for node in fade] == [("fade", "in")])


#============================================
def test_writer_rejects_preexisting_timing_in_plain_or_compatibility_content() -> None:
	"""A writer never merges timing into a plain or compatibility slide tree."""
	for compatibility_wrapped in (False, True):
		presentation = Presentation()
		slide = presentation.slides.add_slide(presentation.slide_layouts[6])
		existing = OxmlElement("p:timing")
		if compatibility_wrapped:
			compatibility = slide._element.makeelement(
				"{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent")
			compatibility.append(existing)
			slide._element.append(compatibility)
		else:
			slide._element.append(existing)
		with pytest.raises(marp_lib.pptx_animation.AnimationError,
				match="without existing timing"):
			marp_lib.pptx_animation.PptxAnimationWriter(slide)
