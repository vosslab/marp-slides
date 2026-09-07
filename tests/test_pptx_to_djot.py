"""Behavioral tests for experimental PPTX-to-extended-Djot conversion."""

# Standard Library
import json
import types
import pathlib
import dataclasses

# PIP3 modules
import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

# Local modules
import marp_lib.importers.pptx_to_djot as pptx_to_djot
import marp_lib.importers.pptx_to_marp as pptx_to_marp
import marp_lib.importers.legacy_slide_plan as legacy_slide_plan
import marp_lib.importers.legacy_djot_emitter as legacy_djot_emitter
import marp_lib.djot_lint
#============================================
def write_png(output_path: pathlib.Path) -> pathlib.Path:
	"""Write one bounded source image."""
	image = Image.new("RGB", (40, 30), (20, 90, 160))
	image.save(output_path)
	return output_path
#============================================
def text_region(
	left: float,
	top: float,
	right: float,
	bottom: float,
	text: str,
	*,
	title_identity: bool = False,
	placeholder_confidence: float = 0.0,
	source_kind: str = "text",
	paragraph_count: int = 1,
) -> legacy_slide_plan.SourceTextRegion:
	"""Build one concise synthetic source text region."""
	return legacy_slide_plan.SourceTextRegion(
		tuple((0, text) for _index in range(paragraph_count)),
		legacy_slide_plan.NormalizedBounds(left, top, right, bottom),
		False,
		1.0 if title_identity else placeholder_confidence,
		title_identity,
		source_kind,
	)
#============================================
def test_conversion_uses_djot_layout_slots_and_omits_notes(tmp_path: pathlib.Path) -> None:
	"""One text-image slide becomes strict-Djot-shaped source without notes."""
	image_path = write_png(tmp_path / "chromosome.png")
	presentation = Presentation()
	slide = presentation.slides.add_slide(presentation.slide_layouts[5])
	assert slide.shapes.title is not None
	slide.shapes.title.text = "Genetics overview"
	text_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.7), Inches(5.2), Inches(3.0))
	text_box.text_frame.text = "Chromosomes carry genes"
	slide.shapes.add_picture(str(image_path), Inches(7.0), Inches(1.5), width=Inches(2.5))
	slide.notes_slide.notes_text_frame.text = "Presenter-only explanation"
	hidden = presentation.slides.add_slide(presentation.slide_layouts[5])
	hidden.element.set("show", "0")
	input_path = tmp_path / "lecture.pptx"
	presentation.save(input_path)
	output_path = tmp_path / "lecture.djot"
	summary = pptx_to_djot.convert_pptx(
		input_path,
		output_path,
		expected_slide_count=2,
		expected_hidden={2},
	)
	djot = output_path.read_text(encoding="utf-8")
	report = json.loads(summary.report_path.read_text(encoding="utf-8"))
	assert summary.visible_slides == 1
	assert summary.hidden_slides == 1
	assert summary.extracted_images == 1
	assert "=== layout: two-panels" in djot
	assert "# Genetics overview" in djot
	assert "@left" in djot and "@right" in djot
	assert "- Chromosomes carry genes" in djot
	assert "![Slide image 1](assets/lecture/image_001.png)" in djot
	assert "Presenter-only explanation" not in djot
	assert "<!--" not in djot
	assert report["hidden_slides"] == [2]
	assert report["slides"][0]["notes_omitted"] == 1
	assert not (summary.report_path.parent / "lecture").exists()
	assert all(not path.is_symlink() for path in summary.report_path.parent.iterdir())
#============================================
def test_empty_notes_part_is_treated_as_no_notes() -> None:
	"""An empty LibreOffice notes relationship never aborts a Djot import."""
	notes_relation = types.SimpleNamespace(reltype="http://example.test/notesSlide")
	slide = types.SimpleNamespace(
		part=types.SimpleNamespace(rels={"notes": notes_relation}),
		notes_slide=types.SimpleNamespace(notes_text_frame=None),
	)
	assert pptx_to_djot.slide_notes(slide) == ()
#============================================
def test_blank_auto_shape_inventory_requires_visible_vector_evidence() -> None:
	"""Only blank auto-shapes with OOXML fill or stroke enter visual inventory."""
	def shape(*, fill: bool = False, stroke: bool = False, xml: str = "<p:sp>") -> object:
		line = types.SimpleNamespace(find=lambda name: object() if stroke and name.endswith("solidFill") else None)
		properties = types.SimpleNamespace(find=lambda name: line if name.endswith("ln") else (
			object() if fill and name.endswith("Fill") else None))
		return types.SimpleNamespace(
			shape_type=MSO_SHAPE_TYPE.AUTO_SHAPE, is_placeholder=False,
			has_text_frame=True, has_table=False,
			text_frame=types.SimpleNamespace(paragraphs=()),
			element=types.SimpleNamespace(xml=xml, find=lambda name: properties if name.endswith("spPr") else None), shape_id=7,
			left=0, top=0, width=20, height=20,
		)
	visible = shape(fill=True)
	stroke = shape(stroke=True, xml='<a:prstGeom prst="line">')
	invisible = shape(xml="<p:txBody><a:rPr><a:solidFill/></a:rPr></p:txBody>")
	assert pptx_to_djot.source_visual_inventory(visible, 100, 100)
	assert pptx_to_djot.source_visual_inventory(stroke, 100, 100)
	assert not pptx_to_djot.source_visual_inventory(invisible, 100, 100)
#============================================
def test_source_text_inventory_preserves_direct_style_role_and_z_order() -> None:
	"""Text planning retains direct source evidence beyond its numeric shape ID."""
	line = types.SimpleNamespace(find=lambda name: object() if name.endswith("solidFill") else None)
	properties = types.SimpleNamespace(find=lambda name: line if name.endswith("ln") else object())
	paragraph = types.SimpleNamespace(level=0, text="Object label", runs=())
	shape = types.SimpleNamespace(
		shape_type=MSO_SHAPE_TYPE.AUTO_SHAPE, is_placeholder=True, has_table=False,
		has_text_frame=True, text_frame=types.SimpleNamespace(paragraphs=(paragraph,)),
		placeholder_format=types.SimpleNamespace(type="OBJECT"),
		element=types.SimpleNamespace(find=lambda name: properties), shape_id=99,
		left=0, top=0, width=20, height=20,
	)
	region = pptx_to_djot.source_text_inventory(shape, None, 100, 100, (2, 1))[0]
	assert (region.has_positive_fill, region.has_positive_line, region.placeholder_role, region.z_order) == (
		True, True, "OBJECT", (2, 1),
	)
#============================================
def test_vector_media_is_validated_and_emf_type_is_normalized() -> None:
	"""Legacy vector blobs retain their actual type only after header validation."""
	wmf_blob = b"\xd7\xcd\xc6\x9a" + b"\x00" * 24
	emf_blob = b"\x01\x00\x00\x00" + b"\x00" * 36 + b" EMF" + b"\x00" * 8

	pptx_to_djot.validate_image_blob(wmf_blob, ".wmf")
	pptx_to_djot.validate_image_blob(emf_blob, ".emf")
	assert pptx_to_djot.image_suffix(emf_blob, ".wmf") == ".emf"
	with pytest.raises(ValueError, match="WMF image header"):
		pptx_to_djot.validate_image_blob(b"not a metafile", ".wmf")
#============================================
def test_dna_text_uses_inline_verbatim_and_ascii_prime_projection() -> None:
	"""Short DNA sequences use the settled source surface without a new delimiter."""
	assert pptx_to_djot.djot_text("5'-AGTACT-3'") == "5&prime;-`AGTACT`-3&prime;"
	assert pptx_to_djot.djot_text("5\u2019-AGTACT-3\u2019") == "5&prime;-`AGTACT`-3&prime;"
	assert pptx_to_djot.djot_text("5\u2032-AGTACT-3\u2032") == "5&prime;-`AGTACT`-3&prime;"
	assert pptx_to_djot.djot_text("Compare ATGC and CGTA") == "Compare `ATGC` and `CGTA`"
#============================================
def test_title_slide_preserves_every_subtitle_line() -> None:
	"""A title slide keeps its multi-line subtitle region as authored H2 lines."""
	slide = pptx_to_marp.SlideData(
		source_index=1,
		hidden=False,
		title_lines=("Course introduction",),
		text_blocks=(pptx_to_marp.TextBlock(0, 0, ((0, "Fall 2026"), (0, "Biology 301")), True),),
		images=(),
		notes=(),
		review_reasons=(),
	)
	lines, layout = pptx_to_djot.render_slide(slide, 100, True)
	assert layout == "title-slide"
	assert lines == [
		"=== layout: title-slide", "", "# Course introduction", "",
		"## Fall 2026", "## Biology 301",
	]
#============================================
def test_converter_refuses_non_djot_or_existing_output(tmp_path: pathlib.Path) -> None:
	"""The experimental importer never overwrites another source format."""
	with pytest.raises(ValueError, match=".djot"):
		pptx_to_djot.validate_output_path(tmp_path / "lecture.md")
	output_path = tmp_path / "lecture.djot"
	output_path.write_text("existing\n", encoding="utf-8")
	with pytest.raises(FileExistsError, match="will not overwrite"):
		pptx_to_djot.validate_output_path(output_path)
#============================================
def test_geometry_beats_misleading_title_placeholder_identity() -> None:
	"""A shallow wide heading wins over an identity-tagged body placeholder."""
	shallow = text_region(
		0.08, 0.06, 0.88, 0.16, "Actual title", placeholder_confidence=1.0,
	)
	misleading = text_region(0.14, 0.44, 0.78, 0.70, "Body placeholder", title_identity=True)
	decision = legacy_slide_plan.select_title((misleading, shallow))
	assert decision.region is shallow
	assert decision.reason == "shallow-wide geometry"
#============================================
def test_unmarked_diagram_label_and_tall_identity_stay_out_of_title_lane() -> None:
	"""Diagram labels, body placeholders, and side titles stay out of the H1 lane."""
	label = text_region(0.08, 0.04, 0.92, 0.18, "Diagram pathway", source_kind="text-box")
	identity = text_region(0.05, 0.25, 0.95, 0.80, "Body", title_identity=True)
	side_title = text_region(0.05, 0.04, 0.40, 0.80, "Side title", title_identity=True)
	decision = legacy_slide_plan.select_title((label, identity))
	solitary_decision = legacy_slide_plan.select_solitary_title((side_title,), ())

	assert decision.region is None
	assert solitary_decision is None
#============================================
def test_solitary_tall_title_identity_becomes_one_global_title() -> None:
	"""One sole text placeholder remains a title despite a tall template box."""
	region = legacy_slide_plan.SourceTextRegion(
		((0, "Course"), (0, "Introduction")),
		legacy_slide_plan.NormalizedBounds(0.05, 0.04, 0.95, 0.88),
		True,
		1.0,
		True,
	)

	image = legacy_slide_plan.SourceImageRegion(
		"assets/overview.png",
		legacy_slide_plan.NormalizedBounds(0.20, 0.35, 0.80, 0.85),
	)
	plan = legacy_slide_plan.plan_slide((region,), (image,))
	centered = legacy_slide_plan.SourceTextRegion(
		((0, "Centered"), (0, "Title")),
		legacy_slide_plan.NormalizedBounds(0.02, 0.20, 0.98, 0.65), True, 1.0, True,
	)
	centered_plan = legacy_slide_plan.plan_slide((centered,), ())
	image_plan = legacy_slide_plan.plan_slide((centered,), (image,))

	assert plan.title.region is region
	assert plan.title.reason == "solitary title identity"
	assert centered_plan.title.region is centered
	assert image_plan.title.region is None
#============================================
def test_geometry_partitions_clear_left_and_right_text_columns() -> None:
	"""Separated text components become named slots in reading order."""
	left = text_region(0.06, 0.28, 0.40, 0.72, "Left explanation", paragraph_count=5)
	right = text_region(0.61, 0.28, 0.94, 0.72, "Right explanation", paragraph_count=5)
	image = legacy_slide_plan.SourceImageRegion(
		"assets/right.png",
		legacy_slide_plan.NormalizedBounds(0.70, 0.30, 0.90, 0.50),
	)

	plan = legacy_slide_plan.plan_slide((left, right), (image,))

	assert tuple(slot.name for slot in plan.slots) == ("left", "right")
	assert plan.slots[1].image_regions == (image,)
#============================================
def test_diagram_plan_retains_paired_internal_heading() -> None:
	"""A qualified diagram keeps nearby narrow and wide annotations together."""
	title = text_region(0.38, 0.12, 0.76, 0.18, "Internal heading", source_kind="auto-shape")
	annotation = text_region(
		0.48, 0.42, 0.56, 0.50, "Embedded diagram heading", source_kind="auto-shape",
	)
	annotation_label = text_region(
		0.70, 0.59, 0.78, 0.67, "Diagram label", source_kind="auto-shape",
	)
	wide_annotation = text_region(
		0.38, 0.66, 0.78, 0.73, "Wide diagram annotation", source_kind="auto-shape",
	)
	overlap_annotation = text_region(
		0.80, 0.64, 0.96, 0.71, "Overlapping annotation", source_kind="auto-shape",
	)
	outside = text_region(0.05, 0.84, 0.42, 0.96, "Editable prose")
	image = legacy_slide_plan.SourceImageRegion(
		"assets/diagram.png",
		legacy_slide_plan.NormalizedBounds(0.35, 0.05, 0.82, 0.75),
	)

	plan = legacy_slide_plan.plan_slide(
		(title, annotation, annotation_label, wide_annotation, overlap_annotation, outside), (image,),
	)

	assert plan.content_region is not None
	assert plan.content_region.text_regions == (
		title, annotation, annotation_label, wide_annotation, overlap_annotation,
	)
	assert plan.title.region is None
	with pytest.raises(ValueError, match="required renderer asset"):
		legacy_slide_plan.require_region_asset(plan.content_region, None)
#============================================
def test_large_title_excluded_diagram_remains_a_component() -> None:
	"""A crossing annotation requests protected-title rendering without cutting pixels."""
	title = legacy_slide_plan.SourceTextRegion(
		((0, "Title"),), legacy_slide_plan.NormalizedBounds(0.05, 0.03, 0.95, 0.13),
		False, 1.0, False, "text", 91,
	)
	first = text_region(0.10, 0.08, 0.18, 0.30, "Diagram label", source_kind="auto-shape")
	second = text_region(0.72, 0.72, 0.80, 0.84, "Other label", source_kind="auto-shape")
	image = legacy_slide_plan.SourceImageRegion(
		"assets/full-slide.png",
		legacy_slide_plan.NormalizedBounds(0.02, 0.02, 0.98, 0.96),
	)
	plan = legacy_slide_plan.plan_slide((title, first, second), (image,))

	assert plan.content_region is not None
	assert plan.content_region.protected_text_shape_ids == (91,)
	assert plan.content_region.bounds.top == image.bounds.top
#============================================
def test_single_auto_shape_caption_keeps_exterior_caption_native() -> None:
	"""One interior auto-shape diagram label leaves an exterior caption editable."""
	image = legacy_slide_plan.SourceImageRegion(
		"assets/diagram.png", legacy_slide_plan.NormalizedBounds(0.05, 0.10, 0.85, 0.80),
	)
	interior = text_region(0.35, 0.40, 0.45, 0.48, "Interior", source_kind="auto-shape")
	exterior = text_region(0.88, 0.40, 0.98, 0.48, "Exterior", source_kind="auto-shape")
	plan = legacy_slide_plan.plan_slide((interior, exterior), (image,))

	assert plan.content_region is not None and plan.content_region.text_regions == (interior,)
	assert plan.slots[0].text_regions == (exterior,)
#============================================
def test_complete_placeholder_list_and_callout_emit_multiple_choice(tmp_path: pathlib.Path) -> None:
	"""Only the complete geometric question structure receives answer semantics."""
	question = legacy_slide_plan.SourceTextRegion(
		((0, "Prompt"), (1, "Choice one"), (1, "Choice two")),
		legacy_slide_plan.NormalizedBounds(0.08, 0.20, 0.82, 0.68), False, 1.0,
		source_ordinal=4,
	)
	answer = legacy_slide_plan.SourceTextRegion(
		((0, "Answer"),), legacy_slide_plan.NormalizedBounds(0.60, 0.63, 0.80, 0.67),
		False, 0.0, source_ordinal=8,
	)
	plan = legacy_slide_plan.plan_slide((question, answer), ())
	top_figure = legacy_slide_plan.SourceImageRegion(
		"assets/figure.png", legacy_slide_plan.NormalizedBounds(0.30, 0.24, 0.70, 0.34),
	)
	non_top_figure = legacy_slide_plan.SourceImageRegion(
		"assets/figure.png", legacy_slide_plan.NormalizedBounds(0.30, 0.50, 0.70, 0.62),
	)
	oversized_figure = legacy_slide_plan.SourceImageRegion(
		"assets/figure.png", legacy_slide_plan.NormalizedBounds(0.08, 0.20, 0.82, 0.62),
	)
	edge_overhang_figure = legacy_slide_plan.SourceImageRegion(
		"assets/figure.png", legacy_slide_plan.NormalizedBounds(0.05, 0.24, 0.45, 0.34),
	)
	flat_prose = legacy_slide_plan.SourceTextRegion(
		((0, "Prompt"), (0, "Repeated prose"), (0, "More prose")), question.bounds,
		False, 1.0, source_ordinal=4,
	)
	planned = pptx_to_djot.PlannedSlide(
		pptx_to_marp.SlideData(4, False, (), (), (), (), ()), plan, visible_page_index=3,
	)

	lines, layout, _reasons = legacy_djot_emitter.render_planned_slide(planned, {}, False)

	assert layout == "multiple-choice"
	assert lines == [
		"=== layout: multiple-choice", "", "@question", "", "Prompt", "", "- Choice one",
		"- Choice two", "", "@answer", "", "Answer",
	]
	assert legacy_slide_plan.plan_slide((flat_prose, answer), ()).multiple_choice is None
	figure_plans = tuple(
		legacy_slide_plan.plan_slide((question, answer), (image,)).multiple_choice
		for image in (top_figure, non_top_figure, oversized_figure, edge_overhang_figure)
	)
	assert figure_plans[0] is not None and figure_plans[0].image is top_figure
	assert figure_plans[1:] == (None, None, None)
	figure_plan = legacy_slide_plan.plan_slide((question, answer), (top_figure,))
	figure_data = pptx_to_marp.SlideData(
		4, False, (), (), (
			pptx_to_marp.ImageAsset(1, 1, 1, 1, "assets/figure.png", "Figure"),
		), (), (),
	)
	figure_lines, _layout, _reasons = legacy_djot_emitter.render_planned_slide(
		pptx_to_djot.PlannedSlide(figure_data, figure_plan, visible_page_index=3), {}, False,
	)
	(tmp_path / "assets").mkdir()
	write_png(tmp_path / "assets" / "figure.png")
	source_path = tmp_path / "deck.djot"
	source_path.write_text("\n".join(figure_lines) + "\n", encoding="utf-8")
	problems, _slides, _images = marp_lib.djot_lint.lint_source(source_path)

	assert figure_lines.count("![Figure](assets/figure.png)") == 1
	assert problems == []


#============================================
def test_regular_text_grid_stays_reviewable_without_table_metadata() -> None:
	"""Text centers alone never claim a faithful native table."""
	regions = tuple(
		text_region(
			0.05 + column * 0.25,
			0.20 + row * 0.18,
			0.20 + column * 0.25,
			0.26 + row * 0.18,
			"cell",
			source_kind="auto-shape",
		)
		for row in range(3)
		for column in range(3)
	)

	plan = legacy_slide_plan.plan_slide(regions, ())

	assert plan.tables == ()
	assert plan.review_reason is not None
#============================================
def test_actual_pptx_table_preserves_blank_cells(tmp_path: pathlib.Path) -> None:
	"""A true PPTX table retains dimensions and an intentional blank cell."""
	presentation = Presentation()
	slide = presentation.slides.add_slide(presentation.slide_layouts[6])
	table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2)).table
	table.first_row = False
	table.cell(0, 0).text = "Header"
	table.cell(1, 1).text = "Value"
	slide.shapes.add_picture(str(write_png(tmp_path / "backing.png")), Inches(0.1), Inches(0.1))
	path = tmp_path / "table.pptx"
	presentation.save(path)
	loaded = Presentation(path)
	loaded_slide = loaded.slides[0]

	plan = pptx_to_djot.plan_source_slide(
		loaded_slide, loaded.slide_width, loaded.slide_height, (),
	)

	assert plan.tables
	assert plan.tables[0].rows[0][1].text_regions == ()
	assert not plan.tables[0].has_header
	assert plan.content_region is None
#============================================
def test_two_actual_pptx_tables_remain_separate(tmp_path: pathlib.Path) -> None:
	"""Independent source tables never merge cells with matching coordinates."""
	presentation = Presentation()
	slide = presentation.slides.add_slide(presentation.slide_layouts[6])
	first = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(2), Inches(1)).table
	second = slide.shapes.add_table(1, 1, Inches(4), Inches(1), Inches(2), Inches(1)).table
	first.cell(0, 0).merge(first.cell(0, 1))
	first.cell(0, 0).text = "First"
	second.cell(0, 0).text = "Second"
	path = tmp_path / "two-tables.pptx"
	presentation.save(path)
	loaded = Presentation(path)

	plan = pptx_to_djot.plan_source_slide(
		loaded.slides[0], loaded.slide_width, loaded.slide_height, (),
	)

	assert len(plan.tables) == 2
	assert plan.tables[0].unsupported_reason is not None
#============================================
def test_side_by_side_text_and_image_use_live_two_panels() -> None:
	"""A live side-by-side topology keeps direct text and image components separate."""
	prose = text_region(.08, .32, .45, .76, "Editable prose")
	image = legacy_slide_plan.SourceImageRegion("assets/photo.png", legacy_slide_plan.NormalizedBounds(.55, .32, .92, .76))

	plan = legacy_slide_plan.plan_slide((prose,), (image,))

	assert tuple(slot.name for slot in plan.slots) == ("left", "right")
	assert not any(slot.flows_in_source_order for slot in plan.slots)


#============================================
def test_genuine_vertical_lane_uses_one_ordered_flow() -> None:
	"""A source-ordered bullet, picture, and caption share one vertical lane."""
	bullet = text_region(.10, .35, .90, .42, "Editable bullet")
	caption = text_region(.10, .75, .90, .80, "Editable caption")
	image = legacy_slide_plan.SourceImageRegion("assets/photo.png", legacy_slide_plan.NormalizedBounds(.22, .40, .78, .70))

	plan = legacy_slide_plan.plan_slide((bullet, caption), (image,))

	assert len(plan.slots) == 1 and plan.slots[0].flows_in_source_order
	assert plan.slots[0].text_regions == (bullet, caption)
	images = (legacy_slide_plan.SourceImageRegion("assets/top.png", legacy_slide_plan.NormalizedBounds(.22, .20, .78, .45)),
		legacy_slide_plan.SourceImageRegion("assets/bottom.png", legacy_slide_plan.NormalizedBounds(.22, .55, .78, .80)))
	image_plan = legacy_slide_plan.plan_slide((), images)
	assert tuple(slot.name for slot in image_plan.slots) == ("top", "bottom")
	assert not any(slot.flows_in_source_order for slot in image_plan.slots)


#============================================
def test_planned_slots_emit_named_cells_without_duplicate_coupled_content() -> None:
	"""The emitter consumes planner slots and reserves coupled labels for one crop."""
	title = text_region(0.05, 0.04, 0.90, 0.14, "Editable title", title_identity=True)
	left = text_region(0.05, 0.30, 0.35, 0.70, "Native left")
	annotation = text_region(0.65, 0.42, 0.72, 0.48, "Diagram label")
	image = legacy_slide_plan.SourceImageRegion(
		"assets/diagram.png", legacy_slide_plan.NormalizedBounds(0.45, 0.25, 0.92, 0.82),
	)
	content = legacy_slide_plan.ContentRegionPlan(
		"content-region-1", legacy_slide_plan.NormalizedBounds(0.45, 0.25, 0.92, 0.82),
		(annotation,), (image,), (91,),
	)
	plan = legacy_slide_plan.LegacySlidePlan(
		legacy_slide_plan.TitleDecision(title, "test"),
		(legacy_slide_plan.SlotPlan("left", (left,)), legacy_slide_plan.SlotPlan("right", ())),
		content,
	)
	data = pptx_to_marp.SlideData(1, False, (), (), (), (), ())
	planned = pptx_to_djot.PlannedSlide(data, plan, (title, left, annotation), (image,), 1)

	lines, layout, _reasons = pptx_to_djot.render_planned_slide(
		planned, {(1, "content-region-1"): "assets/deck/source_region.png"}, False,
	)

	assert layout == "two-panels"
	assert "- Native left" in lines
	assert "![Coupled source region](assets/deck/source_region.png)" in lines
	assert all("Diagram label" not in line for line in lines)
	assert legacy_djot_emitter.render_planned_djot(
		[planned], {(1, "content-region-1"): "assets/deck/source_region.png"},
	)[1][0]["protected_text_shape_ids"] == [91]


#============================================
def test_planned_content_requires_its_renderer_asset() -> None:
	"""A crop plan fails before an emitter can silently flatten its diagram labels."""
	content = legacy_slide_plan.ContentRegionPlan(
		"content-region-1", legacy_slide_plan.NormalizedBounds(0.20, 0.20, 0.80, 0.80), (), (),
	)
	plan = legacy_slide_plan.LegacySlidePlan(
		legacy_slide_plan.TitleDecision(None, "test"),
		(legacy_slide_plan.SlotPlan("body", ()),), content,
	)
	planned = pptx_to_djot.PlannedSlide(
		pptx_to_marp.SlideData(1, False, (), (), (), (), ()), plan, visible_page_index=1,
	)

	with pytest.raises(ValueError, match="required renderer asset"):
		pptx_to_djot.render_planned_slide(planned, {}, False)


#============================================
def test_planned_first_title_slide_retains_all_h2_subtitles() -> None:
	"""Geometry planning does not collapse the authored multi-subtitle title slide."""
	title = text_region(0.08, 0.05, 0.90, 0.14, "Course introduction", title_identity=True)
	subtitle = legacy_slide_plan.SourceTextRegion(
		((0, "Fall 2026"), (0, "Biology 301")),
		legacy_slide_plan.NormalizedBounds(0.20, 0.35, 0.80, 0.55), True, 0.0,
	)
	plan = legacy_slide_plan.LegacySlidePlan(
		legacy_slide_plan.TitleDecision(title, "test"),
		(legacy_slide_plan.SlotPlan("body", (subtitle,)),),
	)
	planned = pptx_to_djot.PlannedSlide(
		pptx_to_marp.SlideData(1, False, (), (), (), (), ()), plan, visible_page_index=1,
	)

	lines, layout, _reasons = pptx_to_djot.render_planned_slide(planned, {}, True)

	assert layout == "title-slide"
	assert lines[-2:] == ["## Fall 2026", "## Biology 301"]


#============================================
def test_separable_images_emit_once_as_atomic_components() -> None:
	"""Each source shape owns one image component, even for a shared asset."""
	first = legacy_slide_plan.SourceImageRegion(
		"assets/shared.png", legacy_slide_plan.NormalizedBounds(0.05, 0.25, 0.40, 0.70),
		source_ordinal=11,
	)
	second = legacy_slide_plan.SourceImageRegion(
		"assets/shared.png", legacy_slide_plan.NormalizedBounds(0.55, 0.25, 0.90, 0.70),
		source_ordinal=12,
	)
	plan = legacy_slide_plan.LegacySlidePlan(
		legacy_slide_plan.TitleDecision(None, "test"),
		(legacy_slide_plan.SlotPlan("body", (), (first, second)),),
	)
	data = pptx_to_marp.SlideData(
		1, False, (), (), (
			pptx_to_marp.ImageAsset(1, 1, 1, 1, "assets/shared.png", "Shared"),
		), (), (),
	)
	planned = pptx_to_djot.PlannedSlide(data, plan, visible_page_index=1)

	components, _reasons = legacy_djot_emitter.emit_components(planned, None)
	assert tuple(component.source_image_ids for component in components) == (
		((11, "assets/shared.png"),), ((12, "assets/shared.png"),),
	)
	assert legacy_djot_emitter.render_planned_slide(planned, {}, False)[1] == "two-panels"


#============================================
def test_unique_image_caption_pair_remains_one_editable_component() -> None:
	"""Tied captions stay atomic while one unique geometric caption pairs."""
	image = legacy_slide_plan.SourceImageRegion(
		"assets/figure.png", legacy_slide_plan.NormalizedBounds(0.25, 0.20, 0.75, 0.50),
		source_ordinal=3,
	)
	caption = legacy_slide_plan.SourceTextRegion(
		((0, "Caption"),), legacy_slide_plan.NormalizedBounds(0.30, 0.54, 0.70, 0.60),
		False, 0.0, False, "auto-shape", 4,
	)
	competing_caption = legacy_slide_plan.SourceTextRegion(
		((0, "Other caption"),), caption.bounds, False, 0.0, False, "auto-shape", 5,
	)
	plan = legacy_slide_plan.LegacySlidePlan(
		legacy_slide_plan.TitleDecision(None, "test"),
		(legacy_slide_plan.SlotPlan("body", (caption, competing_caption), (image,)),),
	)
	planned = pptx_to_djot.PlannedSlide(
		pptx_to_marp.SlideData(1, False, (), (), (
			pptx_to_marp.ImageAsset(1, 1, 1, 1, "assets/figure.png", "Figure"),
		), (), ()), plan,
	)

	tied_components, _reasons = legacy_djot_emitter.emit_components(planned, None)
	unique_plan = legacy_slide_plan.LegacySlidePlan(
		legacy_slide_plan.TitleDecision(None, "test"),
		(legacy_slide_plan.SlotPlan("body", (caption,), (image,)),),
	)
	unique_components, _reasons = legacy_djot_emitter.emit_components(
		dataclasses.replace(planned, plan=unique_plan), None,
	)

	assert len(tied_components) == 3 and all(len(item.lines) == 1 for item in tied_components)
	assert any(component.lines == ("![Figure](assets/figure.png)", "", "- Caption")
		for component in unique_components)


#============================================
def test_caption_pairing_allows_border_contact_but_not_distant_lane_text() -> None:
	"""Captions share a real figure lane, including a small source-border overlap."""
	image = legacy_slide_plan.SourceImageRegion(
		"assets/figure.png", legacy_slide_plan.NormalizedBounds(0.30, 0.20, 0.70, 0.50),
	)
	caption = text_region(0.34, 0.495, 0.66, 0.56, "Caption", source_kind="auto-shape")
	distant = text_region(0.02, 0.495, 0.12, 0.56, "URL", source_kind="auto-shape")

	pairs = legacy_djot_emitter.caption_pairings((caption, distant), (image,))

	assert pairs == {image: caption}
def test_text_flow_keeps_columns_separate_and_groups_matching_footer_lines() -> None:
	"""A changed lane starts a component while same-lane footer text remains editable flow."""
	components = [legacy_djot_emitter.EmissionComponent(
		legacy_slide_plan.NormalizedBounds(left, top, right, bottom), ("- item",), "text",
	) for left, top, right, bottom in (
		(0.08, 0.20, 0.42, 0.62), (0.58, 0.20, 0.92, 0.62),
		(0.08, 0.70, 0.92, 0.76), (0.08, 0.80, 0.92, 0.86),
	)]

	result = legacy_djot_emitter.coalesce_text_flows(components)

	assert len(result) == 3
	assert sum(component.kind == "flow" for component in result) == 1


#============================================
def test_text_flow_keeps_an_interleaved_image_as_its_own_component() -> None:
	"""Flow grouping passes components, rather than their integer indexes, to its guard."""
	components = [
		legacy_djot_emitter.EmissionComponent(
			legacy_slide_plan.NormalizedBounds(0.10, 0.20, 0.40, 0.32), ("- first",), "text",
		),
		legacy_djot_emitter.EmissionComponent(
			legacy_slide_plan.NormalizedBounds(0.10, 0.60, 0.40, 0.72), ("- second",), "text",
		),
		legacy_djot_emitter.EmissionComponent(
			legacy_slide_plan.NormalizedBounds(0.15, 0.40, 0.35, 0.55), ("![image](assets/image.png)",), "image",
		),
	]

	result = legacy_djot_emitter.coalesce_text_flows(components)

	assert len(result) == 3
	assert [component.kind for component in result].count("image") == 1


#============================================
def test_connected_vector_scaffold_consumes_all_labels() -> None:
	"""Connected thin vectors and distributed labels become one renderer component."""
	first = legacy_slide_plan.SourceImageRegion(
		"source-vector-1", legacy_slide_plan.NormalizedBounds(0.30, 0.30, 0.55, 0.34), "vector", 1,
	)
	second = legacy_slide_plan.SourceImageRegion(
		"source-vector-2", legacy_slide_plan.NormalizedBounds(0.51, 0.30, 0.75, 0.34), "vector", 2,
	)
	left = text_region(0.28, 0.38, 0.36, 0.44, "Left", source_kind="auto-shape")
	right = text_region(0.68, 0.38, 0.76, 0.44, "Right", source_kind="auto-shape")

	plan = legacy_slide_plan.plan_slide((left, right), (first, second))
	disconnected = legacy_slide_plan.SourceImageRegion(
		"source-vector-3", legacy_slide_plan.NormalizedBounds(0.05, 0.80, 0.15, 0.84), "vector", 3,
	)

	assert plan.content_region is not None
	assert plan.content_region.image_regions == (first, second)
	disconnected_plan = legacy_slide_plan.plan_slide((left, right), (first, second, disconnected))
	assert disconnected_plan.content_region is not None
	assert disconnected_plan.content_region.image_regions == (first, second)


#============================================
def test_mixed_visual_graph_requires_multiple_coupled_labels() -> None:
	"""One mixed connected visual graph crops, while a photo-caption pair stays native."""
	picture = legacy_slide_plan.SourceImageRegion(
		"assets/diagram.png", legacy_slide_plan.NormalizedBounds(0.20, 0.30, 0.45, 0.55),
	)
	vector = legacy_slide_plan.SourceImageRegion(
		"source-vector-1", legacy_slide_plan.NormalizedBounds(0.42, 0.38, 0.65, 0.42), "vector", 2,
	)
	left = text_region(0.22, 0.58, 0.30, 0.64, "Left", source_kind="auto-shape")
	right = text_region(0.60, 0.46, 0.68, 0.52, "Right", source_kind="auto-shape")
	caption = text_region(0.24, 0.58, 0.42, 0.64, "Caption", source_kind="auto-shape")
	mixed = legacy_slide_plan.plan_slide((left, right), (picture, vector))
	separable = legacy_slide_plan.plan_slide((caption,), (picture,))
	near_vector = legacy_slide_plan.SourceImageRegion(
		"source-vector-2", legacy_slide_plan.NormalizedBounds(0.48, 0.38, 0.65, 0.42), "vector", 3,
	)
	near_only = legacy_slide_plan.plan_slide((left, right), (picture, near_vector))
	container = text_region(0.03, 0.05, 0.97, 0.18, "Container", placeholder_confidence=1.0)
	overlay = text_region(0.28, 0.40, 0.38, 0.47, "URL", source_kind="text-box")
	localized = legacy_slide_plan.plan_slide((container, overlay), (picture,))

	assert mixed.content_region is not None and mixed.content_region.kind == "mixed-visual"
	assert separable.content_region is None and near_only.content_region is None
	assert localized.content_region is not None
	assert localized.content_region.text_regions == (overlay,)
	assert localized.title.region is container


#============================================
def test_table_projection_keeps_blank_cells_editable() -> None:
	"""Supported rectangular source tables become canonical editable pipe tables."""
	filled = text_region(0.10, 0.30, 0.30, 0.40, "Header", source_kind="table")
	blank = legacy_slide_plan.TableCellPlan(0, 1, ())
	table = legacy_slide_plan.TablePlan(
		legacy_slide_plan.NormalizedBounds(0.10, 0.30, 0.70, 0.70), (filled,),
		((legacy_slide_plan.TableCellPlan(0, 0, (filled,)), blank),), 1, 2,
	)
	plan = legacy_slide_plan.LegacySlidePlan(
		legacy_slide_plan.TitleDecision(None, "test"),
		(legacy_slide_plan.SlotPlan("body", ()),), tables=(table,),
	)
	planned = pptx_to_djot.PlannedSlide(
		pptx_to_marp.SlideData(1, False, (), (), (), (), ()), plan, visible_page_index=1,
	)

	lines, _layout, reasons = pptx_to_djot.render_planned_slide(planned, {}, False)

	assert "| Header |  |" in lines
	assert reasons == []


#============================================
def test_merged_table_requires_review_instead_of_flattening() -> None:
	"""A span unsupported by the native table contract fails before publication."""
	table = legacy_slide_plan.TablePlan(
		legacy_slide_plan.NormalizedBounds(0.10, 0.30, 0.70, 0.70), (), (), 1, 1,
		unsupported_reason="merged cells are not representable",
	)
	plan = legacy_slide_plan.LegacySlidePlan(
		legacy_slide_plan.TitleDecision(None, "test"),
		(legacy_slide_plan.SlotPlan("body", ()),), tables=(table,),
	)
	planned = pptx_to_djot.PlannedSlide(
		pptx_to_marp.SlideData(1, False, (), (), (), (), ()), plan, visible_page_index=1,
	)

	with pytest.raises(ValueError, match="source table requires review"):
		pptx_to_djot.render_planned_slide(planned, {}, False)


#============================================
def test_planner_body_and_coupled_component_receive_separate_cells() -> None:
	"""A normal planner body plus diagram crop is emitted as two bounded components."""
	title = text_region(0.05, 0.04, 0.90, 0.14, "Title", title_identity=True)
	prose = text_region(0.05, 0.30, 0.35, 0.75, "Editable prose")
	first = text_region(0.55, 0.35, 0.62, 0.42, "Label one", source_kind="auto-shape")
	second = text_region(0.75, 0.62, 0.82, 0.69, "Label two", source_kind="auto-shape")
	image = legacy_slide_plan.SourceImageRegion(
		"assets/diagram.png", legacy_slide_plan.NormalizedBounds(0.45, 0.25, 0.92, 0.82),
	)
	plan = legacy_slide_plan.plan_slide((title, prose, first, second), (image,))
	planned = pptx_to_djot.PlannedSlide(
		pptx_to_marp.SlideData(1, False, (), (), (), (), ()), plan, visible_page_index=1,
	)

	lines, layout, _reasons = pptx_to_djot.render_planned_slide(
		planned, {(1, "content-region-1"): "assets/deck/crop.png"}, False,
	)

	assert layout in {"two-panels", "stacked-panels"}
	assert "- Editable prose" in lines
	assert "![Coupled source region](assets/deck/crop.png)" in lines


#============================================
def test_publish_rollback_removes_only_just_published_assets(
	tmp_path: pathlib.Path,
	monkeypatch: pytest.MonkeyPatch,
) -> None:
	"""A failed second publication step leaves neither new output destination behind."""
	staging_assets = tmp_path / "staging-assets"
	staging_assets.mkdir()
	(staging_assets / "image.png").write_bytes(b"image")
	staging_djot = tmp_path / "staging.djot"
	staging_djot.write_text("source\n", encoding="utf-8")
	output_path = tmp_path / "deck.djot"

	monkeypatch.setattr(pptx_to_djot.os, "link", lambda _source, _target: (_ for _ in ()).throw(OSError("fail")))
	with pytest.raises(OSError, match="fail"):
		pptx_to_djot.publish_conversion(staging_assets, staging_djot, output_path)

	assert not output_path.exists()
	assert not (tmp_path / "assets" / "deck").exists()


#============================================
def test_pruning_keeps_only_generated_djot_reachable_media(tmp_path: pathlib.Path) -> None:
	"""Consumed source pictures are not published while shared media and crops remain."""
	staging_assets = tmp_path / "staging-assets"
	staging_assets.mkdir()
	for name in ("image_001.png", "shared.png", "source_region_hash.png"):
		(staging_assets / name).write_bytes(b"media")
	staging_djot = tmp_path / "deck.djot"
	staging_djot.write_text(
		"=== layout: two-panels\n\n@left\n![Shared](assets/deck/shared.png)\n\n"
		"@right\n![Coupled source region](assets/deck/source_region_hash.png)\n",
		encoding="utf-8",
	)

	published_count = pptx_to_djot.prune_staged_media(staging_djot, staging_assets, "deck")
	report = {"unique_media_assets": published_count}

	assert not (staging_assets / "image_001.png").exists()
	assert (staging_assets / "shared.png").is_file()
	assert (staging_assets / "source_region_hash.png").is_file()
	assert report["unique_media_assets"] == len(list(staging_assets.iterdir())) == 2


#============================================
def test_pruning_rejects_unsafe_staged_entries(tmp_path: pathlib.Path) -> None:
	"""Private staging fails closed instead of traversing an unexpected asset entry."""
	staging_assets = tmp_path / "staging-assets"
	staging_assets.mkdir()
	(staging_assets / "shared.png").write_bytes(b"media")
	(staging_assets / "nested").mkdir()
	staging_djot = tmp_path / "deck.djot"
	staging_djot.write_text(
		"=== layout: one-panel\n\n@body\n![Shared](assets/deck/shared.png)\n",
		encoding="utf-8",
	)

	with pytest.raises(ValueError, match="must not contain directories"):
		pptx_to_djot.prune_staged_media(staging_djot, staging_assets, "deck")

	assert (staging_assets / "nested").is_dir()


def test_symlinked_assets_parent_is_rejected_without_clobbering(tmp_path: pathlib.Path) -> None:
	"""A destination assets symlink cannot redirect generated importer output."""
	outside = tmp_path / "outside"
	outside.mkdir()
	(tmp_path / "assets").symlink_to(outside, target_is_directory=True)

	with pytest.raises(ValueError, match="real directory"):
		pptx_to_djot.validate_output_path(tmp_path / "deck.djot")

	assert list(outside.iterdir()) == []


#============================================
def test_publish_preserves_output_that_appears_during_asset_commit(
	tmp_path: pathlib.Path,
	monkeypatch: pytest.MonkeyPatch,
) -> None:
	"""A late competing Djot file is never replaced by this conversion transaction."""
	staging_assets = tmp_path / "staging-assets"
	staging_assets.mkdir()
	(staging_assets / "image.png").write_bytes(b"image")
	staging_djot = tmp_path / "staging.djot"
	staging_djot.write_text("new source\n", encoding="utf-8")
	output_path = tmp_path / "deck.djot"
	real_replace = pptx_to_djot.os.replace

	def competing_replace(source: pathlib.Path, target: pathlib.Path) -> None:
		real_replace(source, target)
		output_path.write_text("competing source\n", encoding="utf-8")

	monkeypatch.setattr(pptx_to_djot.os, "replace", competing_replace)
	with pytest.raises(FileExistsError, match="appeared"):
		pptx_to_djot.publish_conversion(staging_assets, staging_djot, output_path)

	assert output_path.read_text(encoding="utf-8") == "competing source\n"
	assert not (tmp_path / "assets" / "deck").exists()


#============================================
@pytest.mark.parametrize(
	("bounds", "expected_layout", "expected_slots"),
	[
		(
			((0.05, 0.20, 0.42, 0.82), (0.58, 0.20, 0.95, 0.82)),
			"two-panels", ("left", "right"),
		),
		(
			((0.08, 0.20, 0.92, 0.45), (0.08, 0.58, 0.92, 0.85)),
			"stacked-panels", ("top", "bottom"),
		),
		(
			((0.05, 0.20, 0.95, 0.42), (0.30, 0.56, 0.70, 0.88)),
			"stacked-panels", ("top", "bottom"),
		),
		(
			((0.05, 0.20, 0.40, 0.45), (0.05, 0.58, 0.40, 0.85), (0.55, 0.20, 0.95, 0.85)),
			"two-plus-one-panels", ("top-left", "bottom-left", "right"),
		),
		(
			((0.05, 0.20, 0.40, 0.90), (0.55, 0.20, 0.85, 0.45), (0.55, 0.55, 0.85, 0.80)),
			"one-plus-two-panels", ("left", "top-right", "bottom-right"),
		),
		(
			((0.05, 0.20, 0.40, 0.45), (0.55, 0.20, 0.90, 0.45), (0.05, 0.60, 0.95, 0.82)),
			"two-over-one-panels", ("top-left", "top-right", "bottom"),
		),
		(
			((0.05, 0.20, 0.35, 0.40), (0.60, 0.20, 0.90, 0.40),
				(0.05, 0.60, 0.35, 0.80), (0.60, 0.60, 0.90, 0.80)),
			"four-panels", ("top-left", "top-right", "bottom-left", "bottom-right"),
		),
		(
			((0.03, 0.20, 0.25, 0.40), (0.39, 0.20, 0.61, 0.40), (0.75, 0.20, 0.97, 0.40),
				(0.03, 0.60, 0.25, 0.80), (0.39, 0.60, 0.61, 0.80), (0.75, 0.60, 0.97, 0.80)),
			"six-panels", ("top-left", "top-center", "top-right", "bottom-left", "bottom-center", "bottom-right"),
		),
	],
)
def test_component_layout_uses_only_unambiguous_supported_topologies(
	bounds: tuple[tuple[float, float, float, float], ...],
	expected_layout: str,
	expected_slots: tuple[str, ...],
) -> None:
	"""Three, four, and six component plans map to their exact registry slots."""
	components = [legacy_djot_emitter.EmissionComponent(
		legacy_slide_plan.NormalizedBounds(*item), ("- item",), "text",
	) for item in bounds]

	layout, slots, _order = legacy_djot_emitter.component_layout(components)

	assert layout == expected_layout
	assert slots == expected_slots


#============================================
@pytest.mark.parametrize(
	"bounds",
	[
		((0.05, 0.20, 0.40, 0.90), (0.55, 0.20, 0.85, 0.55), (0.55, 0.40, 0.85, 0.80)),
	],
)
def test_component_layout_rejects_ambiguous_peer_or_grid_geometry(
	bounds: tuple[tuple[float, float, float, float], ...],
) -> None:
	"""Ambiguous source geometry reports its source slide and visible page."""
	regions = tuple(legacy_slide_plan.SourceTextRegion(
		((0, "item"),), legacy_slide_plan.NormalizedBounds(*item), False, 0.0,
		source_ordinal=index,
	) for index, item in enumerate(bounds))
	plan = legacy_slide_plan.LegacySlidePlan(
		legacy_slide_plan.TitleDecision(None, "test"),
		tuple(legacy_slide_plan.SlotPlan(f"item-{index}", (region,))
			for index, region in enumerate(regions)),
	)
	planned = pptx_to_djot.PlannedSlide(
		pptx_to_marp.SlideData(7, False, (), (), (), (), ()), plan,
		visible_page_index=3,
	)

	with pytest.raises(
		ValueError,
		match=r"source slide 7 \(visible page 3\): overlapping direct components",
	):
		legacy_djot_emitter.render_planned_slide(planned, {}, False)


#============================================
def test_component_layout_rejects_unmatched_three_component_topology() -> None:
	"""Three equal same-row components do not imply an arbitrary native layout."""
	components = [legacy_djot_emitter.EmissionComponent(
		legacy_slide_plan.NormalizedBounds(left, 0.25, right, 0.70), ("- item",), "text",
	) for left, right in ((0.05, 0.27), (0.39, 0.61), (0.73, 0.95))]

	with pytest.raises(ValueError, match="ambiguous topology"):
		legacy_djot_emitter.component_layout(components)


#============================================
def test_border_contact_components_remain_matchable_while_material_overlap_fails() -> None:
	"""Rounding-scale contact preserves topology; spatial collision remains review-only."""
	border_contact = [legacy_djot_emitter.EmissionComponent(
		legacy_slide_plan.NormalizedBounds(left, top, right, bottom), ("- item",), "text",
	) for left, top, right, bottom in ((0.05, 0.20, 0.95, 0.60), (0.05, 0.595, 0.95, 0.95))]
	material_overlap = [legacy_djot_emitter.EmissionComponent(
		legacy_slide_plan.NormalizedBounds(left, top, right, bottom), ("- item",), "text",
	) for left, top, right, bottom in ((0.05, 0.20, 0.55, 0.70), (0.35, 0.45, 0.85, 0.90))]

	assert legacy_djot_emitter.component_layout(border_contact)[0] == "stacked-panels"
	assert legacy_djot_emitter.components_overlap(material_overlap)


#============================================
def test_invalid_source_geometry_fails_closed() -> None:
	"""Out-of-slide bounds never become permissive normalized geometry."""
	with pytest.raises(ValueError, match="outside slide bounds"):
		legacy_slide_plan.normalized_bounds(-20, 0, 5, 20, 100, 100)
