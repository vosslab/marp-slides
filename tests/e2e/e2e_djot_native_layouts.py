#!/usr/bin/env python3
"""Exercise every Djot native layout through PPTX, ODP, and PDF export.

This E2E verifies the source-to-editable-artifact chain.  LibreOffice Impress
click-playback remains the attended presentation-fidelity gate.
"""

# Standard Library
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile
import uuid
import zipfile

# PIP3 modules
import defusedxml.ElementTree
import PIL.Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# Local Modules
import slide_lib.layouts


NAMESPACES = {
	"draw": "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0",
	"office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
	"text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
}


#============================================
def repo_root() -> pathlib.Path:
	"""Return the repository root containing this E2E runner."""
	result = pathlib.Path(__file__).resolve().parents[2]
	return result


#============================================
def require(condition: bool, message: str) -> None:
	"""Raise an actionable failure when an E2E contract is absent."""
	if not condition:
		raise RuntimeError(message)


#============================================
def write_image(image_path: pathlib.Path) -> None:
	"""Create the one deterministic local component image used by gallery."""
	image = PIL.Image.new("RGB", (96, 64), (36, 87, 143))
	image.save(image_path)


#============================================
def title_source(spec: slide_lib.layouts.LayoutSpec) -> list[str]:
	"""Return the permitted global title region for one layout specification."""
	lines: list[str] = []
	if spec.allows_title:
		lines.append(f"# {spec.name} Djot acceptance")
	if spec.allows_subtitle:
		lines.append("## Editable native objects")
	return lines


#============================================
def cell_source(spec: slide_lib.layouts.LayoutSpec, slot_name: str) -> list[str]:
	"""Return minimal valid source for one declared named cell."""
	if spec.name == "multiple-choice" and slot_name == "question":
		return ["Which editable object remains visible?", "", "- Choice A", "- Choice B"]
	if spec.name == "multiple-choice" and slot_name == "answer":
		return ["Answer: Choice A"]
	if spec.name == "gallery":
		return ["![Djot gallery component one](component.png)",
			"![Djot gallery component two](component.png)"]
	return [f"Editable {spec.name} {slot_name}"]


#============================================
def slide_source(spec: slide_lib.layouts.LayoutSpec) -> str:
	"""Build one minimal valid Djot slide from the live layout specification."""
	lines = [f"=== layout: {spec.name}"]
	lines.extend(title_source(spec))
	for slot_name in spec.slot_names:
		lines.append(f"@{slot_name}")
		lines.extend(cell_source(spec, slot_name))
	result = "\n".join(lines)
	return result


#============================================
def write_deck(deck_path: pathlib.Path) -> tuple[str, ...]:
	"""Write one Djot slide per live layout and return their rendering order."""
	layout_names = tuple(slide_lib.layouts.LAYOUTS)
	source = "\n\n".join(slide_source(spec) for spec in slide_lib.layouts.LAYOUTS.values()) + "\n"
	deck_path.write_text(source, encoding="utf-8")
	return layout_names


#============================================
def slide_text(slide: object) -> str:
	"""Return editable text from every PPTX text frame on one slide."""
	result = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
	return result


#============================================
def inspect_multiple_choice_pptx(slide: object) -> None:
	"""Confirm question and answer are separate editable shapes in the final PPTX."""
	question_shapes = [shape for shape in slide.shapes if shape.has_text_frame and
		"Which editable object remains visible?" in shape.text]
	answer_shapes = [shape for shape in slide.shapes if shape.has_text_frame and
		"Answer: Choice A" in shape.text]
	require(len(question_shapes) == 1 and len(answer_shapes) == 1 and
		question_shapes[0] is not answer_shapes[0],
		"multiple-choice retains distinct editable question and answer shapes")


#============================================
def inspect_pptx(pptx_path: pathlib.Path, layout_names: tuple[str, ...]) -> None:
	"""Verify every Djot layout produced editable native PPTX content."""
	presentation = Presentation(pptx_path)
	require(pptx_path.stat().st_size > 0 and len(presentation.slides) == len(layout_names),
		"PPTX exists, is nonempty, and has one slide per live layout")
	for index, layout_name in enumerate(layout_names):
		if layout_name == "blank":
			continue
		text = slide_text(presentation.slides[index])
		require(layout_name in text or layout_name == "multiple-choice",
			f"PPTX {layout_name} retains authored editable text")
	gallery_index = layout_names.index("gallery")
	pictures = [shape for shape in presentation.slides[gallery_index].shapes
		if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
	require(pictures, "PPTX gallery retains the Djot component image as a native picture")
	multiple_choice_index = layout_names.index("multiple-choice")
	inspect_multiple_choice_pptx(presentation.slides[multiple_choice_index])


#============================================
def element_text(element: object) -> str:
	"""Return descendant text from one parsed ODF element."""
	result = "".join(element.itertext())
	return result


#============================================
def editable_odp_objects(page: object) -> list[object]:
	"""Return editable ODP text objects in LibreOffice's supported serializations."""
	frames = [frame for frame in page.findall(".//draw:frame", NAMESPACES)
		if frame.find(".//draw:text-box", NAMESPACES) is not None]
	custom_shapes = [shape for shape in page.findall(".//draw:custom-shape", NAMESPACES)
		if element_text(shape).strip()]
	result = frames + custom_shapes
	return result


#============================================
def inspect_odp(odp_path: pathlib.Path, layout_names: tuple[str, ...]) -> None:
	"""Verify LibreOffice preserved separate editable ODP text and image objects."""
	with zipfile.ZipFile(odp_path) as archive:
		content_root = defusedxml.ElementTree.fromstring(archive.read("content.xml"))
	pages = content_root.findall("./office:body/office:presentation/draw:page", NAMESPACES)
	require(odp_path.stat().st_size > 0 and len(pages) == len(layout_names),
		"ODP exists, is nonempty, and has one page per live layout")
	gallery_index = layout_names.index("gallery")
	gallery_images = pages[gallery_index].findall(".//draw:image", NAMESPACES)
	require(gallery_images, "ODP gallery retains the Djot component image as an editable draw image")
	multiple_choice_index = layout_names.index("multiple-choice")
	multiple_choice_objects = editable_odp_objects(pages[multiple_choice_index])
	question_objects = [item for item in multiple_choice_objects
		if "Which editable object remains visible?" in element_text(item)]
	answer_objects = [item for item in multiple_choice_objects
		if "Answer: Choice A" in element_text(item)]
	require(len(question_objects) == 1 and len(answer_objects) == 1 and
		question_objects[0] is not answer_objects[0],
		"ODP retains multiple-choice question and answer in distinct editable text objects")


#============================================
def inspect_pdf(pdf_path: pathlib.Path, layout_names: tuple[str, ...]) -> None:
	"""Confirm the ODP-derived PDF is nonempty and has one page per layout."""
	result = subprocess.run(["pdfinfo", str(pdf_path)], check=True, capture_output=True, text=True)
	page_match = re.search(r"^Pages:\s+(\d+)\s*$", result.stdout, re.MULTILINE)
	require(pdf_path.stat().st_size > 0 and page_match is not None and
		int(page_match.group(1)) == len(layout_names),
		"ODP-derived PDF is nonempty and has one page per live layout")


#============================================
def run() -> None:
	"""Run source Djot through the public PPTX, ODP, and PDF export chain."""
	root = repo_root()
	output_root = root / "output"
	output_root.mkdir(exist_ok=True)
	stem = f"djot_native_layouts_{uuid.uuid4().hex}"
	workspace = pathlib.Path(tempfile.mkdtemp(prefix=f"{stem}_", dir=output_root))
	deck_path = workspace / f"{stem}.djot"
	pptx_path = output_root / "pptx" / f"{stem}.pptx"
	odp_path = output_root / "odp" / f"{stem}.odp"
	pdf_path = output_root / "pdf" / f"{stem}.pdf"
	try:
		write_image(workspace / "component.png")
		layout_names = write_deck(deck_path)
		command = [sys.executable, "deck_tools.py", "build", str(deck_path), "--format", "pdf"]
		subprocess.run(command, cwd=root, check=True)
		inspect_pptx(pptx_path, layout_names)
		inspect_odp(odp_path, layout_names)
		inspect_pdf(pdf_path, layout_names)
		print("PASS: Djot source retains editable PPTX and ODP semantics through PDF")
	finally:
		for artifact_path in (pptx_path, odp_path, pdf_path):
			artifact_path.unlink(missing_ok=True)
		shutil.rmtree(workspace, ignore_errors=True)


if __name__ == "__main__":
	run()
