"""Fast behavioral tests for bounded source-region rendering helpers."""

# Standard Library
import pathlib
import os

# PIP3 modules
import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

# local repo modules
import slide_lib.importers.source_region_render as source_region_render


#============================================
def test_pixel_bounds_keeps_partial_edges_and_rejects_full_page() -> None:
	"""Normalized crops preserve their right/bottom edge while forbidding full slides."""
	bounds = source_region_render.NormalizedRegion(0.1, 0.2, 0.75, 0.8)
	assert source_region_render.pixel_bounds(bounds, 101, 99) == (10, 19, 76, 80)
	with pytest.raises(source_region_render.SourceRegionRenderError, match="full slide"):
		source_region_render.pixel_bounds(source_region_render.NormalizedRegion(0, 0, 1, 1), 20, 20)


#============================================
def test_request_binds_dense_page_to_its_visible_source_slide() -> None:
	"""A crop's dense page selects the matching source slide across hidden gaps."""
	request = source_region_render.SourceRegionRequest("content-region-1", 3, 2,
		source_region_render.NormalizedRegion(0.1, 0.1, 0.9, 0.9))
	source_region_render.validate_request(request, (1, 3))
	mismatched = source_region_render.SourceRegionRequest("content-region-2", 1, 2,
		source_region_render.NormalizedRegion(0.1, 0.1, 0.9, 0.9))
	with pytest.raises(source_region_render.SourceRegionRenderError, match="does not match"):
		source_region_render.validate_request(mismatched, (1, 3))


#============================================
def test_discover_rendered_pages_accepts_poppler_padding(tmp_path: pathlib.Path) -> None:
	"""Poppler's zero-padded page names retain their numeric page order."""
	for page_name in ("page-01.png", "page-02.png", "page-03.png"):
		(tmp_path / page_name).touch()
	pages = source_region_render.discover_rendered_pages(tmp_path, 3)
	assert [page.name for page in pages] == ["page-01.png", "page-02.png", "page-03.png"]


#============================================
def test_protected_variant_clears_only_requested_top_level_title(tmp_path: pathlib.Path) -> None:
	"""A private normalized PPTX clone clears the title while preserving source and body."""
	source_path = tmp_path / "normalized.pptx"
	presentation = Presentation()
	slide = presentation.slides.add_slide(presentation.slide_layouts[6])
	title = slide.shapes.add_textbox(Inches(1), Inches(0.3), Inches(7), Inches(0.5))
	title.text = "Visible course title"
	body = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(7), Inches(1))
	body.text = "Diagram label remains in the crop"
	presentation.save(source_path)
	authority_bytes = source_path.read_bytes()
	request = source_region_render.SourceRegionRequest("content-region-1", 1, 1,
		source_region_render.NormalizedRegion(0.1, 0.1, 0.9, 0.9), (title.shape_id,))
	clone_path = source_region_render.prepare_protected_variant(source_path, (request,), tmp_path)
	clone_slide = Presentation(clone_path).slides[0]
	assert source_path.read_bytes() == authority_bytes
	assert tuple(shape.text for shape in clone_slide.shapes) == ("", "Diagram label remains in the crop")


#============================================
def test_protected_variant_clears_multiple_structural_text_shapes(tmp_path: pathlib.Path) -> None:
	"""One private crop clone can clear both editable multiple-choice structures."""
	source_path = tmp_path / "normalized.pptx"
	presentation = Presentation()
	slide = presentation.slides.add_slide(presentation.slide_layouts[6])
	question = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
	answer = slide.shapes.add_textbox(Inches(6), Inches(5), Inches(2), Inches(1))
	question.text = "Prompt and choices"
	answer.text = "Answer"
	presentation.save(source_path)
	request = source_region_render.SourceRegionRequest("multiple-choice-visual-1", 1, 1,
		source_region_render.NormalizedRegion(0.1, 0.1, 0.9, 0.9),
		(question.shape_id, answer.shape_id))
	clone = Presentation(source_region_render.prepare_protected_variant(source_path, (request,), tmp_path))
	shapes = {shape.shape_id: shape for shape in clone.slides[0].shapes}

	assert shapes[question.shape_id].text == ""
	assert shapes[answer.shape_id].text == ""


#============================================
def test_protected_variant_rejects_missing_and_grouped_shape_ids(tmp_path: pathlib.Path) -> None:
	"""Only an exact top-level text identity can be protected in a private clone."""
	source_path = tmp_path / "normalized.pptx"
	presentation = Presentation()
	slide = presentation.slides.add_slide(presentation.slide_layouts[6])
	group = slide.shapes.add_group_shape()
	presentation.save(source_path)
	bounds = source_region_render.NormalizedRegion(0.1, 0.1, 0.9, 0.9)
	missing = source_region_render.SourceRegionRequest("content-region-1", 1, 1, bounds, (999,))
	with pytest.raises(source_region_render.SourceRegionRenderError, match="missing"):
		source_region_render.prepare_protected_variant(source_path, (missing,), tmp_path)
	grouped = source_region_render.SourceRegionRequest("content-region-2", 1, 1, bounds,
		(group.shape_id,))
	with pytest.raises(source_region_render.SourceRegionRenderError, match="group"):
		source_region_render.prepare_protected_variant(source_path, (grouped,), tmp_path)


#============================================
def test_page_for_request_routes_protected_crops_to_variant_page(tmp_path: pathlib.Path) -> None:
	"""Protected requests use clone pages while ordinary crops retain source provenance."""
	source_page = tmp_path / "source.png"
	protected_page = tmp_path / "protected.png"
	request = source_region_render.SourceRegionRequest("content-region-1", 1, 1,
		source_region_render.NormalizedRegion(0.1, 0.1, 0.9, 0.9), (8,))
	assert source_region_render.page_for_request(request, [source_page], [protected_page]) == protected_page
	unprotected = source_region_render.SourceRegionRequest("content-region-2", 1, 1,
		source_region_render.NormalizedRegion(0.1, 0.1, 0.9, 0.9))
	assert source_region_render.page_for_request(unprotected, [source_page], [protected_page]) == source_page


#============================================
@pytest.mark.parametrize(("page_names", "reason"), (
	(("page-01.png", "page-02.png", "page-word.png"), "malformed"),
	(("page-01.png", "page-1.png", "page-02.png"), "duplicate"),
	(("page-01.png", "page-03.png"), "exactly"),
))
def test_discover_rendered_pages_rejects_incomplete_or_ambiguous_output(tmp_path: pathlib.Path,
		page_names: tuple[str, ...], reason: str) -> None:
	"""Malformed, duplicate, or missing Poppler pages cannot enter crop rendering."""
	for page_name in page_names:
		(tmp_path / page_name).touch()
	with pytest.raises(source_region_render.SourceRegionRenderError, match=reason):
		source_region_render.discover_rendered_pages(tmp_path, 3)


#============================================
def test_batch_publication_rolls_back_new_assets_after_a_late_failure(tmp_path: pathlib.Path,
		monkeypatch: pytest.MonkeyPatch) -> None:
	"""A later publication error leaves no newly published partial crop asset."""
	staging_dir = tmp_path / "stage"
	staging_dir.mkdir()
	bounds = source_region_render.NormalizedRegion(0.1, 0.1, 0.9, 0.9)
	first = source_region_render.stage_crop(Image.new("RGB", (20, 20), (1, 2, 3)), bounds,
		staging_dir, 0)
	second = source_region_render.stage_crop(Image.new("RGB", (20, 20), (4, 5, 6)), bounds,
		staging_dir, 1)
	real_replace = os.replace
	calls = 0
	def fail_second_replace(source: pathlib.Path, destination: pathlib.Path) -> None:
		nonlocal calls
		calls += 1
		if calls == 2:
			raise OSError("simulated destination failure")
		real_replace(source, destination)
	monkeypatch.setattr(source_region_render.os, "replace", fail_second_replace)
	with pytest.raises(source_region_render.SourceRegionRenderError, match="publication failed"):
		source_region_render.publish_staged_crops((first, second), tmp_path)
	assert not list(tmp_path.glob("source_region_*.png"))
