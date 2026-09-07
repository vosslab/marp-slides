"""Render bounded source-presentation regions as deterministic PNG assets."""

# Standard Library
import os
import re
import shutil
import pathlib
import hashlib
import tempfile
import subprocess
import dataclasses
import zipfile

# PIP3 modules
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# local repo modules
import marp_lib.libreoffice
import marp_lib.importers.odp_to_marp
import marp_lib.importers.pptx_to_marp


PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"
SUPPORTED_SOURCE_SUFFIXES = {".odp", ".pptx"}
MAX_SOURCE_BYTES = 256 * 1024 * 1024
MAX_VISIBLE_PAGES = 500
MAX_RASTER_PIXELS = 100_000_000
MAX_PNG_BYTES = 128 * 1024 * 1024
RENDER_DPI = 144
RENDER_TIMEOUT_SECONDS = 180
ASSET_KEY = re.compile(r"[a-z][a-z0-9_-]{0,127}\Z")
RENDERED_PAGE_NAME = re.compile(r"page-([0-9]+)\.png\Z")


class SourceRegionRenderError(RuntimeError):
	"""Report an actionable bounded source-region rendering failure."""


@dataclasses.dataclass(frozen=True)
class NormalizedRegion:
	"""A validated non-full-slide crop expressed as normalized coordinates."""

	left: float
	top: float
	right: float
	bottom: float


@dataclasses.dataclass(frozen=True)
class SourceRegionRequest:
	"""One geometry-first source crop requested by the legacy slide planner."""

	asset_key: str
	source_slide_index: int
	visible_page_index: int
	bounds: NormalizedRegion
	protected_text_shape_ids: tuple[int, ...] = ()


@dataclasses.dataclass(frozen=True)
class SourceRegionAsset:
	"""One published source-region PNG and the metadata needed by an importer."""

	asset_key: str
	source_slide_index: int
	asset_name: str
	sha256: str
	width: int
	height: int
	dpi: int
	pixel_bounds: tuple[int, int, int, int]


@dataclasses.dataclass(frozen=True)
class StagedCrop:
	"""A fully validated PNG awaiting the batch's final publication phase."""

	temporary_path: pathlib.Path
	asset_name: str
	digest: str
	width: int
	height: int
	pixel_bounds: tuple[int, int, int, int]


#============================================
def validate_source_path(source_path: pathlib.Path) -> None:
	"""Require one bounded, existing source presentation supplied by the importer."""
	# ASVS 2.2.1 and 5.2.1: accept only bounded files of the expected source type.
	if not source_path.is_file() or source_path.suffix.lower() not in SUPPORTED_SOURCE_SUFFIXES:
		raise SourceRegionRenderError("source presentation must be an existing .odp or .pptx file")
	if source_path.stat().st_size > MAX_SOURCE_BYTES:
		raise SourceRegionRenderError("source presentation exceeds the input size limit")
	# ASVS 1.5.1 and 5.2.2: reuse the importer boundary validators before conversion.
	try:
		if source_path.suffix.lower() == ".odp":
			marp_lib.importers.odp_to_marp.validate_odp(source_path)
		else:
			marp_lib.importers.pptx_to_marp.validate_pptx(source_path)
	except (OSError, ValueError, zipfile.BadZipFile) as exc:
		raise SourceRegionRenderError(f"source presentation validation failed: {exc}") from exc


#============================================
def clear_protected_text_shape(shape: object, source_slide_index: int, shape_id: int) -> None:
	"""Clear only text in one verified top-level PPTX shape retained in a private clone."""
	if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
		raise SourceRegionRenderError(
			f"source slide {source_slide_index} protected shape {shape_id} is a group, not a top-level text shape")
	if getattr(shape, "has_table", False) or not getattr(shape, "has_text_frame", False):
		raise SourceRegionRenderError(
			f"source slide {source_slide_index} protected shape {shape_id} is not a text shape")
	for paragraph in shape.text_frame.paragraphs:
		if paragraph.runs:
			for run in paragraph.runs:
				run.text = ""
		elif paragraph.text:
			paragraph.text = ""


#============================================
def prepare_protected_variant(protected_source_path: pathlib.Path,
		requests: tuple[SourceRegionRequest, ...], staging_dir: pathlib.Path) -> pathlib.Path:
	"""Create one validated private PPTX clone with requested protected text cleared."""
	# ASVS 1.5.1, 2.3.3, and 5.3.2: validate input and mutate only an internal clone.
	if protected_source_path.suffix.lower() != ".pptx":
		raise SourceRegionRenderError("protected source presentation must be an existing .pptx file")
	validate_source_path(protected_source_path)
	requested_ids: dict[int, set[int]] = {}
	for request in requests:
		if request.protected_text_shape_ids:
			requested_ids.setdefault(request.source_slide_index, set()).update(request.protected_text_shape_ids)
	if not requested_ids:
		raise SourceRegionRenderError("protected source variant requires protected text shape ids")
	clone_path = staging_dir / "protected_source.pptx"
	shutil.copyfile(protected_source_path, clone_path)
	presentation = Presentation(clone_path)
	for source_slide_index, shape_ids in requested_ids.items():
		if source_slide_index > len(presentation.slides):
			raise SourceRegionRenderError(
				f"source slide {source_slide_index} is absent from the protected source presentation")
		slide = presentation.slides[source_slide_index - 1]
		for shape_id in shape_ids:
			matches = [shape for shape in slide.shapes if shape.shape_id == shape_id]
			if len(matches) != 1:
				raise SourceRegionRenderError(
					f"source slide {source_slide_index} protected shape {shape_id} is missing or duplicated")
			clear_protected_text_shape(matches[0], source_slide_index, shape_id)
	presentation.save(clone_path)
	try:
		marp_lib.importers.pptx_to_marp.validate_pptx(clone_path)
	except (OSError, ValueError, zipfile.BadZipFile) as exc:
		raise SourceRegionRenderError(f"protected source variant validation failed: {exc}") from exc
	return clone_path


#============================================
def validate_region(bounds: NormalizedRegion) -> None:
	"""Reject malformed, empty, and full-page normalized crop requests."""
	# ASVS 2.2.1: positive validation keeps all crop arithmetic within its contract.
	values = (bounds.left, bounds.top, bounds.right, bounds.bottom)
	if not all(isinstance(value, (int, float)) and 0.0 <= value <= 1.0 for value in values):
		raise SourceRegionRenderError("source region bounds must be finite normalized coordinates")
	if bounds.left >= bounds.right or bounds.top >= bounds.bottom:
		raise SourceRegionRenderError("source region bounds must have positive area")
	if bounds.left == 0.0 and bounds.top == 0.0 and bounds.right == 1.0 and bounds.bottom == 1.0:
		raise SourceRegionRenderError("source region may not request a full slide raster")


#============================================
def validate_request(request: SourceRegionRequest, visible_source_indexes: tuple[int, ...]) -> None:
	"""Require one safe asset identity and dense page mapping for a crop request."""
	# ASVS 2.2.1 and 15.2.2: bound identifiers and page work before invoking tools.
	if ASSET_KEY.fullmatch(request.asset_key) is None:
		raise SourceRegionRenderError("source region asset key must be a short lowercase identifier")
	if request.source_slide_index < 1:
		raise SourceRegionRenderError("source slide index must be positive")
	if not 1 <= request.visible_page_index <= len(visible_source_indexes):
		raise SourceRegionRenderError("source region visible page index is outside the visible deck")
	if request.source_slide_index != visible_source_indexes[request.visible_page_index - 1]:
		raise SourceRegionRenderError("source region page does not match its visible source slide")
	if any(not isinstance(shape_id, int) or shape_id < 1 for shape_id in request.protected_text_shape_ids):
		raise SourceRegionRenderError("protected text shape ids must be positive integers")
	if len(set(request.protected_text_shape_ids)) != len(request.protected_text_shape_ids):
		raise SourceRegionRenderError("protected text shape ids must be unique per source region")
	validate_region(request.bounds)


#============================================
def pixel_bounds(bounds: NormalizedRegion, width: int, height: int) -> tuple[int, int, int, int]:
	"""Map normalized bounds to an inclusive-safe Pillow crop rectangle."""
	if width < 1 or height < 1 or width * height > MAX_RASTER_PIXELS:
		raise SourceRegionRenderError("source raster dimensions exceed the image processing limit")
	# Floor left/top and ceil right/bottom without rounding cropped source text inward.
	left = int(bounds.left * width // 1)
	top = int(bounds.top * height // 1)
	right = int(-(-bounds.right * width // 1))
	bottom = int(-(-bounds.bottom * height // 1))
	if left >= right or top >= bottom:
		raise SourceRegionRenderError("source region rounds to an empty pixel crop")
	if left == 0 and top == 0 and right == width and bottom == height:
		raise SourceRegionRenderError("source region rounds to a full slide raster")
	return left, top, right, bottom


#============================================
def require_tool(command_name: str) -> str:
	"""Resolve a fixed Poppler executable without accepting caller command input."""
	command_path = shutil.which(command_name)
	if command_path is None:
		raise SourceRegionRenderError(f"required Poppler tool is unavailable: {command_name}")
	return command_path


#============================================
def run_tool(command: list[str], label: str) -> subprocess.CompletedProcess[str]:
	"""Run one fixed-argv local rendering command with bounded execution time."""
	# ASVS 1.2.5 and 2.3.3: fixed argv and bounded staged work; never invoke a shell.
	try:
		result = subprocess.run(command, check=False, capture_output=True, text=True,
			timeout=RENDER_TIMEOUT_SECONDS)
	except subprocess.TimeoutExpired as exc:
		raise SourceRegionRenderError(f"{label} timed out") from exc
	if result.returncode != 0:
		diagnostic = result.stderr.strip() or result.stdout.strip()
		suffix = f": {diagnostic}" if diagnostic else f" with exit status {result.returncode}"
		raise SourceRegionRenderError(f"{label} failed{suffix}")
	return result


#============================================
def pdf_page_count(pdf_path: pathlib.Path) -> int:
	"""Read the declared PDF page count through fixed Poppler pdfinfo argv."""
	result = run_tool([require_tool("pdfinfo"), str(pdf_path)], "pdfinfo")
	for line in result.stdout.splitlines():
		if line.startswith("Pages:"):
			value = line.partition(":")[2].strip()
			if value.isdecimal() and int(value) > 0:
				return int(value)
	raise SourceRegionRenderError("pdfinfo did not report a positive PDF page count")


#============================================
def render_pdf_pages(pdf_path: pathlib.Path, page_count: int, staging_dir: pathlib.Path) -> list[pathlib.Path]:
	"""Render each PDF page exactly once through a single fixed Poppler invocation."""
	prefix = staging_dir / "page"
	command = [require_tool("pdftoppm"), "-png", "-r", str(RENDER_DPI), "-f", "1",
		"-l", str(page_count), str(pdf_path), str(prefix)]
	run_tool(command, "pdftoppm")
	pages = discover_rendered_pages(staging_dir, page_count)
	return pages


#============================================
def discover_rendered_pages(staging_dir: pathlib.Path, page_count: int) -> list[pathlib.Path]:
	"""Return one exact, numerically ordered Poppler PNG for every requested page."""
	# ASVS 2.2.1 and 5.3.2: parse only internally named files inside private staging.
	if not 1 <= page_count <= MAX_VISIBLE_PAGES:
		raise SourceRegionRenderError("rendered PDF page count is outside the renderer limit")
	pages: dict[int, pathlib.Path] = {}
	for page_path in staging_dir.glob("page-*.png"):
		match = RENDERED_PAGE_NAME.fullmatch(page_path.name)
		if match is None:
			raise SourceRegionRenderError(f"pdftoppm wrote a malformed page name: {page_path.name}")
		if not page_path.is_file():
			raise SourceRegionRenderError(f"pdftoppm page output is not a file: {page_path.name}")
		page_number = int(match.group(1))
		if not 1 <= page_number <= page_count:
			raise SourceRegionRenderError(f"pdftoppm wrote page {page_number} outside the expected range")
		if page_number in pages:
			raise SourceRegionRenderError(f"pdftoppm wrote duplicate page {page_number}")
		pages[page_number] = page_path
	expected = set(range(1, page_count + 1))
	if set(pages) != expected:
		raise SourceRegionRenderError("pdftoppm did not render exactly the expected PDF pages")
	ordered_pages = [pages[number] for number in range(1, page_count + 1)]
	return ordered_pages


#============================================
def decode_png(image_path: pathlib.Path) -> Image.Image:
	"""Decode a bounded real PNG before it becomes an output crop."""
	# ASVS 5.2.2 and 5.2.6: verify magic, decode, and dimensions before use.
	if image_path.stat().st_size > MAX_PNG_BYTES or not image_path.read_bytes().startswith(PNG_SIGNATURE):
		raise SourceRegionRenderError("rendered source page is not a bounded PNG")
	try:
		with Image.open(image_path) as opened_image:
			opened_image.verify()
		with Image.open(image_path) as opened_image:
			if opened_image.width < 1 or opened_image.height < 1 or \
					opened_image.width * opened_image.height > MAX_RASTER_PIXELS:
				raise SourceRegionRenderError("rendered source page exceeds the pixel limit")
			# Check declared dimensions before load so a pixel flood cannot allocate first.
			opened_image.load()
			mode = "RGBA" if "A" in opened_image.getbands() else "RGB"
			decoded = opened_image.convert(mode)
	except (Image.DecompressionBombError, OSError) as exc:
		raise SourceRegionRenderError(f"rendered source page cannot be decoded: {exc}") from exc
	return decoded


#============================================
def stage_crop(image: Image.Image, bounds: NormalizedRegion, staging_dir: pathlib.Path,
		crop_number: int) -> StagedCrop:
	"""Create and validate one crop without changing the caller's assets directory."""
	pixels = pixel_bounds(bounds, image.width, image.height)
	crop_width = pixels[2] - pixels[0]
	crop_height = pixels[3] - pixels[1]
	if crop_width * crop_height > MAX_RASTER_PIXELS:
		raise SourceRegionRenderError("source region crop exceeds the pixel limit")
	cropped = image.crop(pixels)
	temporary_path = staging_dir / f"crop_{crop_number:03d}.png"
	cropped.save(temporary_path, format="PNG", optimize=False)
	blob = temporary_path.read_bytes()
	if len(blob) > MAX_PNG_BYTES or not blob.startswith(PNG_SIGNATURE):
		raise SourceRegionRenderError("source region crop exceeds the PNG output limit")
	with Image.open(temporary_path) as verified:
		verified.verify()
	digest = hashlib.sha256(blob).hexdigest()
	asset_name = f"source_region_{digest}.png"
	return StagedCrop(temporary_path, asset_name, digest, cropped.width, cropped.height, pixels)


#============================================
def publish_staged_crops(crops: tuple[StagedCrop, ...], assets_dir: pathlib.Path) -> None:
	"""Atomically publish a validated crop batch or remove only this call's new assets."""
	created: list[pathlib.Path] = []
	try:
		for crop in crops:
			destination = assets_dir / crop.asset_name
			if not destination.resolve().is_relative_to(assets_dir.resolve()):
				raise SourceRegionRenderError("source region destination escapes the assets directory")
			if destination.exists():
				if hashlib.sha256(destination.read_bytes()).hexdigest() != crop.digest:
					raise SourceRegionRenderError("source region digest destination has conflicting content")
				continue
			# ASVS 2.3.3 and 5.3.2: publish only validated staged data at internal destinations.
			os.replace(crop.temporary_path, destination)
			created.append(destination)
	except (OSError, SourceRegionRenderError) as exc:
		for destination in created:
			destination.unlink(missing_ok=True)
		raise SourceRegionRenderError(f"source region publication failed: {exc}") from exc


#============================================
def render_presentation_pages(source_path: pathlib.Path, visible_slide_count: int,
		staging_dir: pathlib.Path) -> list[pathlib.Path]:
	"""Convert and rasterize one validated presentation once inside private staging."""
	try:
		pdf_path = marp_lib.libreoffice.convert_file(source_path, staging_dir, "pdf",
			timeout_seconds=RENDER_TIMEOUT_SECONDS)
	except marp_lib.libreoffice.LibreOfficeError as exc:
		raise SourceRegionRenderError(f"LibreOffice PDF conversion failed: {exc}") from exc
	page_count = pdf_page_count(pdf_path)
	if page_count != visible_slide_count:
		raise SourceRegionRenderError(
			f"source PDF has {page_count} pages but importer mapped {visible_slide_count} visible slides")
	pages = render_pdf_pages(pdf_path, page_count, staging_dir)
	return pages


#============================================
def page_for_request(request: SourceRegionRequest, source_pages: list[pathlib.Path],
		protected_pages: list[pathlib.Path] | None) -> pathlib.Path:
	"""Select the original or text-protected raster page for one requested content crop."""
	pages = protected_pages if request.protected_text_shape_ids else source_pages
	if pages is None:
		raise SourceRegionRenderError(
			f"source slide {request.source_slide_index} requires a protected source render variant")
	return pages[request.visible_page_index - 1]


#============================================
def render_source_regions(source_path: pathlib.Path, assets_dir: pathlib.Path,
		visible_source_indexes: tuple[int, ...], requests: tuple[SourceRegionRequest, ...],
		protected_source_path: pathlib.Path | None = None) -> tuple[SourceRegionAsset, ...]:
	"""Render requested bounded source regions after one source-to-PDF conversion."""
	validate_source_path(source_path)
	if not assets_dir.is_dir():
		raise SourceRegionRenderError("assets directory must already exist")
	if not 1 <= len(visible_source_indexes) <= MAX_VISIBLE_PAGES:
		raise SourceRegionRenderError("visible slide count is outside the renderer limit")
	if any(source_index < 1 for source_index in visible_source_indexes) or \
			len(set(visible_source_indexes)) != len(visible_source_indexes):
		raise SourceRegionRenderError("visible source indexes must be unique positive values")
	for request in requests:
		validate_request(request, visible_source_indexes)
	if len({request.asset_key for request in requests}) != len(requests):
		raise SourceRegionRenderError("source region asset keys must be unique")
	if not requests:
		return ()
	# ASVS 2.3.3: all intermediate files live below the caller-owned assets directory.
	with tempfile.TemporaryDirectory(prefix=".source_region_", dir=assets_dir) as temporary_name:
		staging_dir = pathlib.Path(temporary_name)
		source_render_dir = staging_dir / "source"
		source_render_dir.mkdir()
		pages = render_presentation_pages(source_path, len(visible_source_indexes), source_render_dir)
		protected_requests = tuple(request for request in requests if request.protected_text_shape_ids)
		protected_pages: list[pathlib.Path] | None = None
		if protected_requests:
			if protected_source_path is None:
				raise SourceRegionRenderError("protected source path is required for protected source regions")
			protected_render_dir = staging_dir / "protected"
			protected_render_dir.mkdir()
			clone_path = prepare_protected_variant(protected_source_path, protected_requests,
				protected_render_dir)
			protected_pages = render_presentation_pages(clone_path, len(visible_source_indexes),
				protected_render_dir)
		crops: list[StagedCrop] = []
		for crop_number, request in enumerate(requests):
			image = decode_png(page_for_request(request, pages, protected_pages))
			crops.append(stage_crop(image, request.bounds, staging_dir, crop_number))
		publish_staged_crops(tuple(crops), assets_dir)
		assets = [SourceRegionAsset(request.asset_key, request.source_slide_index,
			crop.asset_name, crop.digest, crop.width, crop.height, RENDER_DPI, crop.pixel_bounds)
			for request, crop in zip(requests, crops, strict=True)]
	return tuple(assets)
