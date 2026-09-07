"""Convert a trusted PPTX with the bounded native PPTX-to-Djot importer."""

# Standard Library
import argparse
import dataclasses
import hashlib
import json
import math
import os
import pathlib
import re
import shutil
import stat
import tempfile

# PIP3 modules
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

# Local modules
import marp_lib.djot_lint as djot_lint
import marp_lib.djot_parser
import marp_lib.importers.legacy_geometry as legacy_geometry
import marp_lib.importers.legacy_djot_emitter as legacy_djot_emitter
import marp_lib.importers.legacy_slide_plan as legacy_slide_plan
import marp_lib.importers.pptx_to_marp as pptx_common
import marp_lib.importers.source_region_render as source_region_render


WMF_HEADERS = (b"\xd7\xcd\xc6\x9a", b"\x01\x00\x09\x00\x00\x03", b"\x02\x00\x09\x00\x00\x03")
EMF_SIGNATURE_OFFSET = 40
EMF_SIGNATURE = b" EMF"
SUPPORTED_IMAGE_SUFFIXES = {".emf", ".gif", ".jpg", ".png", ".wmf"}
ASSET_NAME = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.-]*$")
PRIMED_SEQUENCE = re.compile(
	r"([35])[\u2032\u2019'']-([ACGTU][ACGTU|/,.]{2,}[ACGTU])-[\u2032\u2019'']([35])"
)
PRIME_MARK = re.compile(r"([35])[\u2032\u2019'']")
DNA_SEQUENCE = re.compile(r"(?<![A-Za-z0-9`])([ACGTU][ACGTU|/,.]{2,}[ACGTU])(?![A-Za-z0-9`])")


@dataclasses.dataclass(frozen=True)
class ConversionSummary:
	"""User-facing counts for one Djot conversion."""

	visible_slides: int
	editable_slides: int
	hidden_slides: int
	extracted_images: int
	review_slides: int
	output_path: pathlib.Path
	report_path: pathlib.Path


#============================================
def djot_text(text: str) -> str:
	"""Normalize short DNA notation into settled Djot authoring forms."""
	text = PRIMED_SEQUENCE.sub(r"\1&prime;-`\2`-\3&prime;", text)
	text = PRIME_MARK.sub(r"\1&prime;", text)
	return DNA_SEQUENCE.sub(r"`\1`", text)


#============================================
def validate_image_blob(blob: bytes, suffix: str) -> None:
	"""Validate one raster, WMF, or EMF image before writing it."""
	# ASVS 5.2.2 and 5.2.6: validate the declared type and bounded content.
	if suffix not in SUPPORTED_IMAGE_SUFFIXES:
		raise ValueError(f"unsupported PPTX image type: {suffix}")
	if len(blob) > pptx_common.MAX_MEMBER_BYTES:
		raise ValueError("PPTX image exceeds the per-image size limit")
	if suffix == ".wmf":
		# ASVS 5.2.2: accept only recognized WMF headers.
		if not blob.startswith(WMF_HEADERS):
			raise ValueError("invalid WMF image header")
		return
	if suffix == ".emf":
		# ASVS 5.2.2: validate an EMF record type and its required signature.
		if blob[:4] != b"\x01\x00\x00\x00" or blob[EMF_SIGNATURE_OFFSET:44] != EMF_SIGNATURE:
			raise ValueError("invalid EMF image header")
		return
	pptx_common.validate_image_blob(blob, suffix)


#============================================
def image_suffix(blob: bytes, declared_suffix: str) -> str:
	"""Correct LibreOffice's EMF-as-WMF metadata before validation."""
	if declared_suffix == ".wmf" and blob[EMF_SIGNATURE_OFFSET:44] == EMF_SIGNATURE:
		return ".emf"
	return declared_suffix


#============================================
def image_asset(
	shape: object,
	assets_dir: pathlib.Path,
	djot_root: pathlib.PurePosixPath,
	known_images: dict[str, str],
) -> pptx_common.ImageAsset:
	"""Extract one picture into a content-addressed source asset."""
	blob = shape.image.blob
	suffix = f".{shape.image.ext.lower()}"
	if suffix == ".jpeg":
		suffix = ".jpg"
	suffix = image_suffix(blob, suffix)
	validate_image_blob(blob, suffix)
	digest = hashlib.sha256(blob).hexdigest()
	asset_name = known_images.get(digest)
	if asset_name is None:
		asset_name = f"image_{len(known_images) + 1:03d}{suffix}"
		# ASVS 5.3.2: an archive name never selects the output destination.
		(assets_dir / asset_name).write_bytes(blob)
		known_images[digest] = asset_name
	djot_path = (djot_root / asset_name).as_posix()
	asset_number = int(re.search(r"\d+", asset_name).group())
	return pptx_common.ImageAsset(
		left=shape.left,
		top=shape.top,
		width=shape.width,
		height=shape.height,
		markdown_path=djot_path,
		alt_text=f"Slide image {asset_number}",
	)


#============================================
def shape_inventory(
	shape: object,
	title_id: int | None,
	assets_dir: pathlib.Path,
	djot_root: pathlib.PurePosixPath,
	known_images: dict[str, str],
) -> tuple[list[str], list[pptx_common.TextBlock], list[pptx_common.ImageAsset], list[str]]:
	"""Extract supported semantic objects recursively from one shape."""
	titles: list[str] = []
	text_blocks: list[pptx_common.TextBlock] = []
	images: list[pptx_common.ImageAsset] = []
	review_reasons: list[str] = []
	if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
		for child in shape.shapes:
			child_parts = shape_inventory(child, title_id, assets_dir, djot_root, known_images)
			titles.extend(child_parts[0])
			text_blocks.extend(child_parts[1])
			images.extend(child_parts[2])
			review_reasons.extend(child_parts[3])
		return titles, text_blocks, images, review_reasons
	if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
		images.append(image_asset(shape, assets_dir, djot_root, known_images))
		return titles, text_blocks, images, review_reasons
	if getattr(shape, "has_table", False):
		for row in shape.table.rows:
			cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
			if cells:
				line = " | ".join(djot_text(pptx_common.markdown_text(cell)) for cell in cells)
				text_blocks.append(pptx_common.TextBlock(shape.left, shape.top, ((0, line),), False))
		return titles, text_blocks, images, review_reasons
	if getattr(shape, "has_text_frame", False):
		line_values: list[tuple[int, str]] = []
		for paragraph in shape.text_frame.paragraphs:
			text = djot_text(pptx_common.paragraph_markdown(paragraph))
			if text:
				line_values.append((paragraph.level, text))
		lines = tuple(line_values)
		if not lines:
			if shape.shape_type != MSO_SHAPE_TYPE.PLACEHOLDER:
				review_reasons.append(f"ignored non-content shape type {shape.shape_type}")
			return titles, text_blocks, images, review_reasons
		if shape.shape_id == title_id:
			titles.extend(text for _level, text in lines)
		else:
			text_blocks.append(
				pptx_common.TextBlock(
					shape.left,
					shape.top,
					lines,
					pptx_common.is_subtitle_shape(shape),
				)
			)
		return titles, text_blocks, images, review_reasons
	if shape.shape_type != MSO_SHAPE_TYPE.PLACEHOLDER:
		review_reasons.append(f"ignored non-content shape type {shape.shape_type}")
	return titles, text_blocks, images, review_reasons


#============================================
def source_text_inventory(
	shape: object,
	title_id: int | None,
	slide_width: int,
	slide_height: int,
	z_order: tuple[int, ...] = (),
) -> list[legacy_slide_plan.SourceTextRegion]:
	"""Read every text shape into geometry-first planning evidence."""
	regions: list[legacy_slide_plan.SourceTextRegion] = []
	if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
		for child_index, child in enumerate(shape.shapes):
			child_z_order = (*z_order, child_index)
			regions.extend(source_text_inventory(
				child, title_id, slide_width, slide_height, child_z_order,
			))
		return regions
	has_positive_fill, has_positive_line = shape_style_evidence(shape)
	placeholder_role = shape_placeholder_role(shape)
	if getattr(shape, "has_table", False):
		bounds = legacy_geometry.normalized_bounds(
			shape.left, shape.top, shape.width, shape.height, slide_width, slide_height,
		)
		row_count = len(shape.table.rows)
		column_count = len(shape.table.columns)
		has_header = bool(shape.table.first_row)
		merged = any(
			cell.is_spanned or (cell.is_merge_origin and (
				cell.span_height > 1 or cell.span_width > 1
			))
			for row in shape.table.rows
			for cell in row.cells
		)
		unsupported_reason = "merged source table cells require span-aware emission" if merged else None
		for row_index, row in enumerate(shape.table.rows):
			for column_index, cell in enumerate(row.cells):
				text = cell.text.strip()
				if text:
					regions.append(
						legacy_slide_plan.SourceTextRegion(
							((0, djot_text(pptx_common.markdown_text(text))),),
							bounds,
							False,
							0.0,
							False,
							"table",
							shape.shape_id * 10_000 + row_index * column_count + column_index,
							row_index,
							column_index,
							row_count,
							column_count,
							shape.shape_id,
							has_header,
							unsupported_reason,
							rotation_degrees=0.0,
							has_positive_fill=has_positive_fill,
							has_positive_line=has_positive_line,
							placeholder_role=placeholder_role,
							z_order=z_order,
						)
					)
		return regions
	if not getattr(shape, "has_text_frame", False):
		return regions
	paragraphs: list[tuple[int, str]] = []
	for paragraph in shape.text_frame.paragraphs:
		text = djot_text(pptx_common.paragraph_markdown(paragraph))
		if text:
			paragraphs.append((paragraph.level, text))
	if not paragraphs:
		return regions
	bounds = legacy_geometry.normalized_bounds(
		shape.left,
		shape.top,
		shape.width,
		shape.height,
		slide_width,
		slide_height,
	)
	confidence = 1.0 if shape.is_placeholder else 0.0
	if shape.is_placeholder:
		source_kind = "text"
	elif shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
		source_kind = "text-box"
	else:
		source_kind = "auto-shape"
	regions.append(
		legacy_slide_plan.SourceTextRegion(
			tuple(paragraphs),
			bounds,
			pptx_common.is_subtitle_shape(shape),
			confidence,
			shape.shape_id == title_id,
			source_kind,
			shape.shape_id,
			rotation_degrees=getattr(shape, "rotation", 0.0),
			has_positive_fill=has_positive_fill,
			has_positive_line=has_positive_line,
			placeholder_role=placeholder_role,
			z_order=z_order,
		)
	)
	return regions


#============================================
def source_text_regions(
	slide: object,
	slide_width: int,
	slide_height: int,
) -> tuple[legacy_slide_plan.SourceTextRegion, ...]:
	"""Collect source text regions before title or slot interpretation."""
	title_id = pptx_common.title_shape_id(slide)
	regions: list[legacy_slide_plan.SourceTextRegion] = []
	for shape_index, shape in enumerate(slide.shapes):
		regions.extend(source_text_inventory(
			shape, title_id, slide_width, slide_height, (shape_index,),
		))
	ordered = tuple(sorted(
		regions, key=lambda item: (item.bounds.top, item.bounds.left, item.source_ordinal),
	))
	return ordered


#============================================
def source_visual_inventory(
	shape: object,
	slide_width: int,
	slide_height: int,
	z_order: tuple[int, ...] = (),
) -> list[legacy_slide_plan.SourceImageRegion]:
	"""Collect picture and non-text vector anchors for spatial planning."""
	regions: list[legacy_slide_plan.SourceImageRegion] = []
	if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
		for child_index, child in enumerate(shape.shapes):
			regions.extend(source_visual_inventory(child, slide_width, slide_height, (*z_order, child_index)))
		return regions
	if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
		kind = "picture"
	elif (
		shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
		and not getattr(shape, "is_placeholder", False)
		and not getattr(shape, "has_table", False)
		and (
			not getattr(shape, "has_text_frame", False)
			or not any(paragraph.text.strip() for paragraph in shape.text_frame.paragraphs)
		)
		and shape_has_visible_vector_content(shape)
	):
		kind = "connector" if shape_is_stroked_connector(shape) else "vector"
	elif not getattr(shape, "has_text_frame", False) and not getattr(shape, "has_table", False):
		kind = "connector" if shape_is_stroked_connector(shape) else "vector"
	else:
		return regions
	line = None
	if shape.width <= 0 or shape.height <= 0:
		if kind != "connector":
			return regions
		line = degenerate_connector(shape, slide_width, slide_height)
		if line is None:
			return regions
		bounds = line.footprint
	else:
		bounds = legacy_geometry.normalized_bounds(
			shape.left, shape.top, shape.width, shape.height, slide_width, slide_height,
		)
	regions.append(
		legacy_slide_plan.SourceImageRegion(
			f"source-{kind}-{shape.shape_id}",
			bounds,
			kind,
			shape.shape_id,
			line,
			z_order,
		)
	)
	return regions


#============================================
def shape_style_evidence(shape: object) -> tuple[bool, bool]:
	"""Return direct positive fill and line evidence from one source shape."""
	properties = shape.element.find(qn("p:spPr"))
	if properties is None:
		return False, False
	fill_names = ("a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill")
	visible_fill = any(properties.find(qn(name)) is not None for name in fill_names)
	line = properties.find(qn("a:ln"))
	visible_line = line is not None and any(line.find(qn(name)) is not None for name in fill_names)
	return visible_fill, visible_line


#============================================
def shape_placeholder_role(shape: object) -> str | None:
	"""Return a stable placeholder role independent of its source shape ID."""
	if not shape.is_placeholder:
		return None
	role = str(shape.placeholder_format.type)
	return role.split()[0].upper()


#============================================
def shape_has_visible_vector_content(shape: object) -> bool:
	"""Require direct shape-property fill or stroke evidence for blank vectors."""
	visible_fill, visible_line = shape_style_evidence(shape)
	return visible_fill or visible_line


#============================================
def shape_is_stroked_connector(shape: object) -> bool:
	"""Retain only OOXML-evidenced connector geometry as a separate visual role."""
	xml = shape.element.xml
	return shape_has_visible_vector_content(shape) and (
		"<p:cxnSp" in xml or 'prst="line"' in xml or "<a:headEnd" in xml or "<a:tailEnd" in xml
	)


#============================================
def degenerate_connector(
	shape: object, slide_width: int, slide_height: int,
) -> legacy_geometry.DegenerateConnectorFootprint | None:
	"""Return one planning-only footprint for a positively stroked line shape."""
	properties = shape.element.find(qn("p:spPr"))
	if properties is None:
		return None
	line = properties.find(qn("a:ln"))
	width = None if line is None else line.get("w")
	if width is None or not width.isdigit():
		return None
	stroke_width = float(width)
	if not math.isfinite(stroke_width) or stroke_width <= 0:
		return None
	return legacy_geometry.degenerate_connector_footprint(
		shape.left, shape.top, shape.width, shape.height, stroke_width, slide_width, slide_height,
	)


#============================================
def source_visual_regions(
	slide: object,
	slide_width: int,
	slide_height: int,
) -> tuple[legacy_slide_plan.SourceImageRegion, ...]:
	"""Return source visuals in z-order for geometry-only planning."""
	visuals: list[legacy_slide_plan.SourceImageRegion] = []
	for shape_index, shape in enumerate(slide.shapes):
		visuals.extend(source_visual_inventory(shape, slide_width, slide_height, (shape_index,)))
	return tuple(visuals)


#============================================
def plan_source_slide(
	slide: object,
	slide_width: int,
	slide_height: int,
	images: tuple[pptx_common.ImageAsset, ...],
) -> legacy_slide_plan.LegacySlidePlan:
	"""Expose a renderer-neutral import plan for one validated PPTX slide."""
	text_regions = source_text_regions(slide, slide_width, slide_height)
	visual_regions = source_visual_regions(slide, slide_width, slide_height)
	picture_sources: dict[legacy_geometry.NormalizedBounds, list[legacy_slide_plan.SourceImageRegion]] = {}
	for region in visual_regions:
		if region.source_kind == "picture":
			picture_sources.setdefault(region.bounds, []).append(region)
	extracted_images: list[legacy_slide_plan.SourceImageRegion] = []
	for image in images:
		bounds = legacy_geometry.normalized_bounds(
			image.left, image.top, image.width, image.height, slide_width, slide_height,
		)
		sources = picture_sources.get(bounds, [])
		source = sources.pop(0) if sources else None
		extracted_images.append(legacy_slide_plan.SourceImageRegion(
			image.markdown_path, bounds, "picture", 0 if source is None else source.source_ordinal,
			z_order=() if source is None else source.z_order,
		))
	known_bounds = {image.bounds for image in extracted_images}
	image_regions = tuple(extracted_images) + tuple(
		image for image in visual_regions if image.bounds not in known_bounds
	)
	plan = legacy_slide_plan.plan_slide(text_regions, image_regions)
	return plan


#============================================
def plan_text_blocks(
	plan: legacy_slide_plan.LegacySlidePlan,
	legacy_blocks: list[pptx_common.TextBlock],
	regions: tuple[legacy_slide_plan.SourceTextRegion, ...],
	slide_width: int,
	slide_height: int,
) -> tuple[tuple[str, ...], list[pptx_common.TextBlock]]:
	"""Project a geometry-selected title while retaining unsupported table text."""
	if plan.title.region is None:
		return (), legacy_blocks
	title_lines = tuple(text for _level, text in plan.title.region.paragraphs)
	region_blocks = [
		pptx_common.TextBlock(
			int(region.bounds.left * slide_width),
			int(region.bounds.top * slide_height),
			region.paragraphs,
			region.is_subtitle,
		)
		for region in regions
		if region is not plan.title.region
	]
	for block in legacy_blocks:
		# Unmatched table blocks remain editable until native-table planning selects them.
		matches_region = any(
			block.lines == region_block.lines
			and abs(block.left - region_block.left) <= 1
			and abs(block.top - region_block.top) <= 1
			for region_block in region_blocks
		)
		if not matches_region:
			region_blocks.append(block)
	return title_lines, region_blocks


#============================================
def slide_notes(slide: object) -> tuple[str, ...]:
	"""Read source notes safely, but never render them as slide source."""
	has_notes_part = any(rel.reltype.endswith("/notesSlide") for rel in slide.part.rels.values())
	if not has_notes_part:
		return ()
	notes_frame = slide.notes_slide.notes_text_frame
	if notes_frame is None:
		return ()
	return tuple(line.strip() for line in notes_frame.text.splitlines() if line.strip())


#============================================
def extract_slides(
	presentation: object,
	assets_dir: pathlib.Path,
	djot_root: pathlib.PurePosixPath,
	expected_hidden: set[int] | None,
) -> tuple[list[legacy_djot_emitter.PlannedSlide], int]:
	"""Extract slides while preserving source order and authoritative visibility."""
	known_images: dict[str, str] = {}
	slides: list[legacy_djot_emitter.PlannedSlide] = []
	visible_page_index = 0
	for source_index, slide in enumerate(presentation.slides, start=1):
		pptx_hidden = slide.element.get("show") == "0"
		hidden = pptx_hidden if expected_hidden is None else source_index in expected_hidden
		if expected_hidden is not None and pptx_hidden != hidden:
			raise RuntimeError(f"ODP and PPTX visibility disagree on slide {source_index}")
		if hidden:
			slides.append(legacy_djot_emitter.PlannedSlide(
			pptx_common.SlideData(source_index, True, (), (), (), (), ()), None,
		))
			continue
		visible_page_index += 1
		titles: list[str] = []
		text_blocks: list[pptx_common.TextBlock] = []
		images: list[pptx_common.ImageAsset] = []
		review_reasons: list[str] = []
		title_id = pptx_common.title_shape_id(slide)
		for shape in slide.shapes:
			parts = shape_inventory(shape, title_id, assets_dir, djot_root, known_images)
			titles.extend(parts[0])
			text_blocks.extend(parts[1])
			images.extend(parts[2])
			review_reasons.extend(parts[3])
		regions = source_text_regions(slide, presentation.slide_width, presentation.slide_height)
		plan = plan_source_slide(
			slide,
			presentation.slide_width,
			presentation.slide_height,
			tuple(images),
		)
		planned_titles, text_blocks = plan_text_blocks(
			plan,
			text_blocks,
			regions,
			presentation.slide_width,
			presentation.slide_height,
		)
		if planned_titles:
			titles = list(planned_titles)
		line_count = sum(len(block.lines) for block in text_blocks)
		character_count = sum(
			len(text) for block in text_blocks for _level, text in block.lines
		)
		if line_count > 10 or character_count > 1200:
			review_reasons.append("dense text requires post-conversion polish")
		text_blocks.sort(key=lambda block: (block.top, block.left))
		images.sort(key=lambda image: (image.top, image.left))
		data = pptx_common.SlideData(
				source_index=source_index,
				hidden=False,
				title_lines=tuple(titles),
				text_blocks=tuple(text_blocks),
				images=tuple(images),
				notes=slide_notes(slide),
				review_reasons=tuple(sorted(set(review_reasons))),
		)
		slides.append(legacy_djot_emitter.PlannedSlide(
			data, plan, regions, source_visual_regions(
				slide, presentation.slide_width, presentation.slide_height,
			), visible_page_index,
		))
	return slides, len(known_images)


#============================================
def body_lines(slide: pptx_common.SlideData) -> list[str]:
	"""Render positioned text blocks as Djot lists and subtitles."""
	lines: list[str] = []
	for block in slide.text_blocks:
		for level, text in block.lines:
			if block.is_subtitle:
				lines.append(f"## {text}")
			else:
				lines.append(f"{'  ' * level}- {text}")
	return lines


#============================================
def titled_slide(slide: pptx_common.SlideData) -> tuple[list[str], list[str]]:
	"""Return source title lines and remaining body content."""
	texts = body_lines(slide)
	title_lines = list(slide.title_lines)
	if not title_lines and texts:
		title_lines = [texts.pop(0).removeprefix("- ")]
	if not title_lines:
		title_lines = [f"Slide {slide.source_index}"]
	return [f"# {' '.join(title_lines)}"], texts


#============================================
def render_slide(
	slide: pptx_common.SlideData,
	slide_width: int,
	is_first: bool,
) -> tuple[list[str], str]:
	"""Map one semantic slide to supported native layouts and Djot slots."""
	heading_lines, texts = titled_slide(slide)
	images = list(slide.images)
	if not images:
		if is_first and (not texts or all(text.startswith("## ") for text in texts)):
			return ["=== layout: title-slide", "", *heading_lines, "", *texts], "title-slide"
		if not texts:
			return ["=== layout: title-only", "", *heading_lines], "title-only"
		return ["=== layout: one-panel", "", *heading_lines, "", "@body", "", *texts], "one-panel"
	if len(images) == 1 and not texts:
		return [
			"=== layout: one-panel", "", *heading_lines, "", "@body", "",
			legacy_djot_emitter.image_djot(images[0]),
		], "one-panel"
	if len(images) == 1:
		image = images[0]
		center = (image.left + image.width / 2) / slide_width
		if center >= 0.55:
			return [
				"=== layout: two-panels", "", *heading_lines, "", "@left", "", *texts,
				"", "@right", "", legacy_djot_emitter.image_djot(image),
			], "two-panels"
		if center <= 0.45:
			return [
				"=== layout: two-panels", "", *heading_lines, "", "@left", "",
				legacy_djot_emitter.image_djot(image), "", "@right", "", *texts,
			], "two-panels"
	image_lines = [legacy_djot_emitter.image_djot(image) for image in images]
	if not texts:
		return ["=== layout: gallery", "", *heading_lines, "", "@gallery", "", *image_lines], "gallery"
	return [
		"=== layout: two-panels", "", *heading_lines, "", "@left", "", *texts,
		"", "@right", "", *image_lines,
	], "two-panels"


#============================================
def render_djot(
	slides: list[pptx_common.SlideData],
	slide_width: int,
) -> tuple[str, list[dict[str, object]]]:
	"""Render visible slide source and an auditable per-slide import report."""
	visible_slides = [slide for slide in slides if not slide.hidden]
	djot_lines: list[str] = []
	records: list[dict[str, object]] = []
	for visible_index, slide in enumerate(visible_slides):
		if visible_index:
			djot_lines.append("")
		slide_lines, layout = render_slide(slide, slide_width, visible_index == 0)
		djot_lines.extend(slide_lines)
		records.append(
			{
				"source_slide": slide.source_index,
				"layout": layout,
				"text_blocks": len(slide.text_blocks),
				"images": len(slide.images),
				"notes_omitted": len(slide.notes),
				"review_reasons": list(slide.review_reasons),
			}
		)
	return "\n".join(djot_lines).rstrip() + "\n", records


#============================================
def validate_output_path(output_path: pathlib.Path) -> None:
	"""Protect an established Djot destination from replacement."""
	if output_path.suffix.lower() != ".djot":
		raise ValueError("output must use the .djot extension")
	if output_path.exists():
		raise FileExistsError("output Djot source already exists; import will not overwrite it")
	assets_parent = output_path.parent / "assets"
	if assets_parent.is_symlink() or (assets_parent.exists() and not assets_parent.is_dir()):
		raise ValueError("output assets parent must be a real directory")
	asset_path = assets_parent / output_path.stem
	if asset_path.exists() or asset_path.is_symlink():
		raise FileExistsError("output asset directory already exists")


#============================================
def validate_staged_djot(staging_djot: pathlib.Path, staging_assets: pathlib.Path,
		deck_stem: str) -> None:
	"""Parse, lay out, and resolve local staged assets before publication."""
	asset_target = staging_djot.parent / "assets" / deck_stem
	asset_target.parent.mkdir(exist_ok=True)
	os.symlink(staging_assets, asset_target)
	try:
		problems, _slides, _images = djot_lint.lint_source(staging_djot)
		if problems:
			problem = problems[0]
			raise ValueError(
				f"generated Djot staging validation failed: {problem.path}:{problem.line}: {problem.message}"
			)
	finally:
		asset_target.unlink(missing_ok=True)


#============================================
def staged_media_references(staging_djot: pathlib.Path, deck_stem: str) -> set[str]:
	"""Return validated direct-child media names from parsed generated Djot."""
	deck = marp_lib.djot_parser.parse_deck(staging_djot)
	expected_parent = pathlib.PurePosixPath("assets") / deck_stem
	references: set[str] = set()
	for slide in deck.slides:
		images = list(djot_lint.images_in_blocks(slide.blocks))
		for cell in slide.cells:
			images.extend(djot_lint.images_in_blocks(cell.blocks))
		for image in images:
			path = pathlib.PurePosixPath(image.source)
			if path.parent != expected_parent or len(path.parts) != 3:
				raise ValueError("generated Djot image must use its direct local asset namespace")
			if not ASSET_NAME.fullmatch(path.name) or path.suffix.lower() not in SUPPORTED_IMAGE_SUFFIXES:
				raise ValueError("generated Djot image has an unsafe local asset name")
			references.add(path.name)
	return references


#============================================
def prune_staged_media(staging_djot: pathlib.Path, staging_assets: pathlib.Path,
		deck_stem: str) -> int:
	"""Keep only parsed-Djot-reachable regular media in private staging assets."""
	references = staged_media_references(staging_djot, deck_stem)
	media_names: set[str] = set()
	for candidate in staging_assets.iterdir():
		metadata = candidate.lstat()
		if stat.S_ISLNK(metadata.st_mode):
			raise ValueError("staged assets must not contain symlinks")
		if stat.S_ISDIR(metadata.st_mode):
			raise ValueError("staged assets must not contain directories")
		if not stat.S_ISREG(metadata.st_mode):
			raise ValueError("staged assets must contain only regular media files")
		if not ASSET_NAME.fullmatch(candidate.name) or candidate.suffix.lower() not in SUPPORTED_IMAGE_SUFFIXES:
			raise ValueError("staged assets contain an unsafe media filename")
		media_names.add(candidate.name)
	if missing := references - media_names:
		raise ValueError(f"generated Djot references missing staged media: {sorted(missing)[0]}")
	for candidate in staging_assets.iterdir():
		if candidate.name not in references:
			candidate.unlink()
	return len(references)


#============================================
def publish_conversion(staging_assets: pathlib.Path, staging_djot: pathlib.Path,
		output_path: pathlib.Path) -> None:
	"""Publish both new destinations or remove only this call's partial assets."""
	assets_parent = output_path.parent / "assets"
	if assets_parent.is_symlink() or (assets_parent.exists() and not assets_parent.is_dir()):
		raise ValueError("output assets parent must be a real directory")
	assets_parent.mkdir(exist_ok=True)
	final_assets = assets_parent / output_path.stem
	if output_path.exists() or output_path.is_symlink() or final_assets.exists() or final_assets.is_symlink():
		raise FileExistsError("output destination appeared during conversion; nothing was overwritten")
	if not final_assets.parent.resolve().is_relative_to(output_path.parent.resolve()):
		raise ValueError("output assets directory escapes the output parent")
	published_assets = False
	try:
		os.replace(staging_assets, final_assets)
		published_assets = True
		if output_path.exists() or output_path.is_symlink():
			raise FileExistsError("output Djot source appeared during conversion")
		os.link(staging_djot, output_path)
	except OSError:
		if published_assets:
			shutil.rmtree(final_assets)
		raise


#============================================
def convert_pptx(
	input_path: pathlib.Path,
	output_path: pathlib.Path,
	*,
	expected_slide_count: int | None = None,
	expected_hidden: set[int] | None = None,
	source_name: str | None = None,
	render_source_path: pathlib.Path | None = None,
) -> ConversionSummary:
	"""Convert one trusted PPTX into supported bounded Djot source."""
	input_path = input_path.resolve()
	output_path = output_path.resolve()
	# ASVS 1.5.2: validate the existing OOXML boundary before python-pptx reads it.
	pptx_common.validate_pptx(input_path)
	validate_output_path(output_path)
	presentation = Presentation(input_path)
	if expected_slide_count is not None and len(presentation.slides) != expected_slide_count:
		raise RuntimeError("ODP and normalized PPTX slide counts disagree")
	output_path.parent.mkdir(parents=True, exist_ok=True)
	with tempfile.TemporaryDirectory(prefix=".pptx_to_djot_", dir=output_path.parent) as temporary_name:
		temporary_root = pathlib.Path(temporary_name)
		staging_assets = temporary_root / "assets"
		staging_assets.mkdir()
		djot_root = pathlib.PurePosixPath("assets") / output_path.stem
		slides, image_count = extract_slides(presentation, staging_assets, djot_root, expected_hidden)
		visible_slides = [slide for slide in slides if not slide.data.hidden]
		if not visible_slides:
			raise ValueError("presentation contains no visible slides")
		visible_indexes = tuple(slide.data.source_index for slide in visible_slides)
		requests: list[source_region_render.SourceRegionRequest] = []
		request_details: dict[str, tuple[legacy_djot_emitter.PlannedSlide, str]] = {}
		for planned in visible_slides:
			regions = [] if planned.plan is None else [planned.plan.content_region]
			if planned.plan is not None and planned.plan.multiple_choice is not None:
				regions.append(planned.plan.multiple_choice.question_visual_region)
			for content in regions:
				if content is None:
					continue
				global_key = f"source-{planned.data.source_index}-{content.asset_key}"
				requests.append(source_region_render.SourceRegionRequest(
					global_key, planned.data.source_index, planned.visible_page_index or 0,
					source_region_render.NormalizedRegion(
						content.bounds.left, content.bounds.top, content.bounds.right, content.bounds.bottom,
					),
					content.protected_text_shape_ids,
				))
				request_details[global_key] = (planned, content.asset_key)
		render_source = (render_source_path or input_path).resolve()
		rendered = source_region_render.render_source_regions(
			render_source, staging_assets, visible_indexes, tuple(requests),
			protected_source_path=input_path,
		)
		assets: dict[tuple[int, str], str] = {}
		asset_details: dict[tuple[int, str], source_region_render.SourceRegionAsset] = {}
		for asset in rendered:
			planned, local_key = request_details[asset.asset_key]
			key = (planned.data.source_index, local_key)
			assets[key] = (djot_root / asset.asset_name).as_posix()
			asset_details[key] = asset
		# Fail before publication if a plan requires a missing crop.
		for planned in visible_slides:
			regions = [] if planned.plan is None else [planned.plan.content_region]
			if planned.plan is not None and planned.plan.multiple_choice is not None:
				regions.append(planned.plan.multiple_choice.question_visual_region)
			for content in regions:
				legacy_slide_plan.require_region_asset(
					content, None if content is None or \
						(planned.data.source_index, content.asset_key) not in assets else {
						content.asset_key: assets.get((planned.data.source_index, content.asset_key), ""),
					},
				)
		djot, records = legacy_djot_emitter.render_planned_djot(slides, assets)
		for record, planned in zip(records, visible_slides, strict=True):
			content = planned.plan.content_region if planned.plan else None
			if planned.plan is not None:
				record["omitted_vectors"] = [
					{"source_ordinal": item.source_ordinal, "bounds": dataclasses.asdict(item.bounds),
						"reason": "decorative-vector"}
					for item in planned.plan.omitted_vectors
				]
				record["review_vectors"] = [
					{"source_ordinal": item.source_ordinal, "bounds": dataclasses.asdict(item.bounds),
						"reason": "unconsumed-vector"}
					for item in planned.plan.review_vectors
				]
				choice = planned.plan.multiple_choice
				if choice is not None and choice.question_visual_region is not None:
					visual = choice.question_visual_region
					key = (planned.data.source_index, visual.asset_key)
					asset = asset_details[key]
					record["multiple_choice"]["question_visual_region"] = {
						"local_key": visual.asset_key,
						"global_key": asset.asset_key,
						"asset": assets[key],
						"sha256": asset.sha256,
						"dpi": asset.dpi,
						"normalized_bounds": dataclasses.asdict(visual.bounds),
						"pixel_bounds": list(asset.pixel_bounds),
						"included_source_image_refs": [item.asset_reference for item in visual.image_regions],
						"included_source_ordinals": [item.source_ordinal for item in visual.image_regions],
						"overlay_text_count": len(visual.text_regions),
						"protected_text_shape_ids": list(visual.protected_text_shape_ids),
						"classification_reason": visual.classification_reason,
					}
			if content is None:
				continue
			key = (planned.data.source_index, content.asset_key)
			asset = asset_details[key]
			record["content_region"] = {
				"local_key": content.asset_key,
				"global_key": asset.asset_key,
				"asset": assets[key],
				"sha256": asset.sha256,
				"dpi": asset.dpi,
				"normalized_bounds": dataclasses.asdict(content.bounds),
				"pixel_bounds": list(asset.pixel_bounds),
				"included_source_image_refs": [item.asset_reference for item in content.image_regions],
				"included_source_ordinals": [item.source_ordinal for item in content.image_regions],
				"coupled_text_count": len(content.text_regions),
				"protected_text_shape_ids": list(content.protected_text_shape_ids),
				"kind": content.kind,
				"classification_reason": content.classification_reason,
			}
		staging_djot = temporary_root / output_path.name
		staging_djot.write_text(djot, encoding="utf-8")
		published_media_count = prune_staged_media(staging_djot, staging_assets, output_path.stem)
		report = {
			"source": source_name or input_path.name,
			"slide_count": len(slides),
			"visible_slides": len(visible_slides),
			"hidden_slides": [slide.data.source_index for slide in slides if slide.data.hidden],
			"unique_media_assets": published_media_count,
			"slides": records,
		}
		report_path = staging_assets / "import_report.json"
		report_path.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
		validate_staged_djot(staging_djot, staging_assets, output_path.stem)
		publish_conversion(staging_assets, staging_djot, output_path)
	final_report = output_path.parent / "assets" / output_path.stem / "import_report.json"
	review_count = sum(bool(record["review_reasons"]) for record in records)
	return ConversionSummary(
		visible_slides=len(visible_slides),
		editable_slides=len(visible_slides),
		hidden_slides=len(slides) - len(visible_slides),
		extracted_images=json.loads(final_report.read_text(encoding="utf-8"))["unique_media_assets"],
		review_slides=review_count,
		output_path=output_path,
		report_path=final_report,
	)


#============================================
def parse_args() -> argparse.Namespace:
	"""Parse standalone PPTX-to-Djot importer arguments."""
	parser = argparse.ArgumentParser(description=__doc__)
	parser.add_argument("input_file", type=pathlib.Path, help="trusted source PPTX")
	parser.add_argument("--output", type=pathlib.Path, help="new extended-Djot source path")
	return parser.parse_args()


#============================================
def main() -> None:
	"""Run one standalone PPTX-to-Djot conversion."""
	args = parse_args()
	output_path = args.output or args.input_file.with_suffix(".djot")
	summary = convert_pptx(args.input_file, output_path)
	print(
		f"Converted {summary.visible_slides} visible slides: "
		f"{summary.editable_slides} editable, {summary.review_slides} layout review, "
		f"{summary.hidden_slides} hidden, {summary.extracted_images} content images"
	)
	print(f"Djot: {summary.output_path}")
	print(f"Import report: {summary.report_path}")


if __name__ == "__main__":
	main()
